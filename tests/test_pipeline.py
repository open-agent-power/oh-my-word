import copy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.shared import Mm, Pt

from word_skills.common import read_json, write_json
from word_skills.delivery import build, review_init, finalize
from word_skills.common import digest
from word_skills.evidence import contract, validate
from word_skills.ingest import extract, ingest
from word_skills.render import default_document, fingerprint, render
from word_skills.inspect import inspect_docx


@pytest.fixture
def sample(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("系统采用 Java 17。\n系统提供查询接口。\n", encoding="utf-8")
    case = tmp_path / "case"
    ingest([source], case)
    write_json(case / "evidence.json", {"schema_version": 1, "claims": [{
        "id": "C1", "statement": "系统采用 Java 17。", "source_id": "S001", "locators": ["L1"],
        "quote": "系统采用 Java 17。", "status": "verified", "verified_by": "fixture", "verified_at": "2026-10-15"}]})
    draft = {"schema_version": 1, "mode": "technical-report", "title": "系统技术报告", "sections": [
        {"id": k, "title": v, "blocks": [{"id": k + "b", "type": "paragraph", "intent": "fact", "text": "系统采用 Java 17。", "claim_ids": ["C1"]}]}
        for k, v in contract("technical-report")["required_sections"].items()]}
    path = tmp_path / "draft.json"
    write_json(path, draft)
    return source, case, draft, path


def codes(result):
    return {x["code"] for x in result["issues"]}


def test_grounded_draft_passes(sample):
    assert validate(sample[2], sample[1])["passed"]


def test_fake_qps_cannot_be_laundered_by_unrelated_citation(sample):
    sample[2]["sections"][0]["blocks"][0]["text"] += "吞吐量达到 10000 QPS。"
    assert "UNSUPPORTED_NUMBER" in codes(validate(sample[2], sample[1]))


def test_fake_claim_number_fails_even_if_marked_verified(sample):
    ledger = read_json(sample[1] / "evidence.json")
    ledger["claims"][0]["statement"] = "吞吐量为 10000 QPS。"
    write_json(sample[1] / "evidence.json", ledger)
    assert "CLAIM_NUMBER" in codes(validate(sample[2], sample[1]))


@pytest.mark.parametrize("field,value,code", [("quote", "并不存在", "QUOTE"), ("locators", ["L2"], "QUOTE"), ("status", "unverified", "UNVERIFIED"), ("source_id", "missing", "CLAIM_SOURCE")])
def test_bad_evidence(sample, field, value, code):
    ledger = read_json(sample[1] / "evidence.json")
    ledger["claims"][0][field] = value
    write_json(sample[1] / "evidence.json", ledger)
    assert code in codes(validate(sample[2], sample[1]))


def test_original_change_invalidates(sample):
    sample[0].write_text("changed", encoding="utf-8")
    assert "LIVE_SOURCE_CHANGED" in codes(validate(sample[2], sample[1]))


def test_snapshot_change_invalidates(sample):
    (sample[1] / "sources/S001/original.txt").write_text("changed", encoding="utf-8")
    assert "SOURCE_CHANGED" in codes(validate(sample[2], sample[1]))


def test_extraction_change_invalidates(sample):
    (sample[1] / "sources/S001/extracted.md").write_text("changed", encoding="utf-8")
    assert "SOURCE_CHANGED" in codes(validate(sample[2], sample[1]))


def test_missing_section(sample):
    sample[2]["sections"].pop()
    assert "MISSING_SECTION" in codes(validate(sample[2], sample[1]))


def test_heading_jump(sample):
    sample[2]["sections"][0]["level"] = 3
    assert "HEADING_LEVEL" in codes(validate(sample[2], sample[1]))


def test_table_number_is_checked(sample):
    sample[2]["sections"][0]["blocks"] = [{"id": "table", "type": "table", "intent": "analysis", "headers": ["QPS", "结论"], "rows": [["10000", "通过"]], "claim_ids": []}]
    assert "UNSUPPORTED_NUMBER" in codes(validate(sample[2], sample[1]))


def test_missing_evidence_prevents_docx(sample, tmp_path):
    sample[2]["sections"][0]["blocks"][0]["claim_ids"] = []
    write_json(sample[3], sample[2])
    output = tmp_path / "out"
    assert build(sample[3], sample[1], output, no_preview=True)["state"] == "blocked"
    assert not (output / "output.docx").exists()


def test_no_preview_remains_draft(sample, tmp_path):
    output = tmp_path / "out"
    result = build(sample[3], sample[1], output, no_preview=True)
    assert result["state"] == "draft"
    assert Document(output / "output.docx").paragraphs
    with pytest.raises(ValueError):
        review_init(output)


def test_refuses_output_overwrite(sample, tmp_path):
    with pytest.raises(ValueError, match="exists"):
        build(sample[3], sample[1], tmp_path)


def test_reference_style_and_margins(sample, tmp_path):
    template = default_document()
    template.sections[0].left_margin = Mm(30)
    template.styles["Heading 1"].font.size = Pt(18)
    template.add_paragraph("reference body is replaced only in style mode")
    path = tmp_path / "reference.docx"
    template.save(path)
    output = tmp_path / "out.docx"
    render(sample[2], output, path)
    before, after = fingerprint(path), fingerprint(output)
    assert before["sections"] == after["sections"]
    assert before["styles"] == after["styles"]
    assert all("reference body" not in p.text for p in Document(output).paragraphs)


def test_docxtpl_slots(sample, tmp_path):
    template = default_document()
    template.add_paragraph("{{ title }}", "Title")
    template.add_paragraph("{{p body }}")
    path = tmp_path / "slots.docx"
    template.save(path)
    out = tmp_path / "filled.docx"
    render(sample[2], out, path, "slots")
    doc = Document(out)
    assert doc.paragraphs[0].text == sample[2]["title"]
    assert any(p.text == "系统采用 Java 17。" for p in doc.paragraphs)


def test_template_expressions_rejected(sample, tmp_path):
    doc = default_document()
    doc.add_paragraph("{{ title.__class__ }}")
    doc.add_paragraph("{{p body }}")
    path = tmp_path / "slots.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="Slot template"):
        render(sample[2], tmp_path / "out.docx", path, "slots")


def test_macro_payload_rejected(tmp_path):
    path = tmp_path / "macro.docx"
    default_document().save(path)
    with ZipFile(path, "a", ZIP_DEFLATED) as z:
        z.writestr("word/vbaProject.bin", b"not executed")
    with pytest.raises(ValueError, match="Macro"):
        extract(path)


def test_ingest_word_records_structural_locations(tmp_path):
    path = tmp_path / "source.docx"
    doc = default_document()
    doc.add_paragraph("正文证据")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格证据"
    doc.save(path)
    text, segments, notices = extract(path)
    assert "正文证据" in text
    assert any(s["text"] == "表格证据" and "/row/" in s["locator"] for s in segments)
    assert notices


def test_case_inside_input_refused(tmp_path):
    with pytest.raises(ValueError, match="outside"):
        ingest([tmp_path], tmp_path / "case")


def test_two_modes_different_contracts():
    assert set(contract("technical-report")["required_sections"]) != set(contract("design-spec")["required_sections"])


def test_long_title_and_table_style(sample, tmp_path):
    draft = copy.deepcopy(sample[2])
    draft["title"] = "面向跨部门协作与长期维护场景的技术资料追溯和文档生成系统技术报告"
    draft["sections"].append({"id": "extra", "title": "附加说明", "blocks": [{"id": "t", "type": "table", "intent": "analysis", "headers": ["项目", "说明"], "rows": [["边界", "只记录已知事实，不制造实验结论。" * 8]], "widths": [1, 4], "style": "Light Shading Accent 1", "claim_ids": []}]})
    assert validate(draft, sample[1])["passed"]
    render(draft, tmp_path / "long.docx")
    assert len(Document(tmp_path / "long.docx").tables) == 1


@pytest.mark.parametrize("body,title", [("prefix {{p body }} suffix", "{{ title }}"), ("{{p body }}", "")])
def test_slot_contract_rejects_lossy_inputs(sample, tmp_path, body, title):
    doc = default_document()
    if title:
        doc.add_paragraph(title, "Title")
    doc.add_paragraph(body)
    template = tmp_path / "template.docx"
    doc.save(template)
    with pytest.raises(ValueError, match="standalone"):
        render(sample[2], tmp_path / "out.docx", template, "slots")


def test_slot_tables_use_template_width(sample, tmp_path):
    doc = default_document()
    doc.sections[0].left_margin = doc.sections[0].right_margin = Mm(40)
    doc.add_paragraph("{{ title }}", "Title")
    doc.add_paragraph("{{p body }}")
    template = tmp_path / "template.docx"
    doc.save(template)
    sample[2]["sections"][0]["blocks"].append({"id": "table", "type": "table", "intent": "analysis", "headers": ["项目", "说明"], "rows": [["测试", "长内容" * 20]], "claim_ids": []})
    out = tmp_path / "out.docx"
    render(sample[2], out, template, "slots")
    assert not inspect_docx(out, sample[2])


def test_missing_title_detected(sample, tmp_path):
    out = tmp_path / "out.docx"
    render(sample[2], out)
    doc = Document(out)
    doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element)
    doc.save(out)
    assert "DOCX_TITLE" in {i["code"] for i in inspect_docx(out, sample[2])}


@pytest.fixture
def reviewable(sample, tmp_path, monkeypatch):
    # Unit fixture only: not evidence that the real PDF rendering path works.
    sample[2]["task_requirements"] = [{"id": "T1", "text": "说明技术栈", "acceptance": "明确语言版本及依据", "block_ids": ["backgroundb"]}]
    write_json(sample[3], sample[2])
    def fake_preview(docx, out, *args):
        (out / "preview.pdf").write_bytes(b"unit-preview")
        (out / "page-1.png").write_bytes(b"unit-page")
        return {"backend": "unit-fixture", "page_images": ["page-1.png"], "pages": 1}
    monkeypatch.setattr("word_skills.delivery.preview", fake_preview)
    monkeypatch.setattr("word_skills.delivery.inspect_pdf", lambda *a: {"issues": [], "pages": [1]})
    out = tmp_path / "reviewable"
    assert build(sample[3], sample[1], out)["state"] == "awaiting-review"
    r = review_init(out)
    r.update(reviewer="unit fixture", reviewed_at="2026-10-15")
    r["content"] = {"status": "pass", "notes": "test fixture"}
    for row in r["task_coverage"]:
        row.update(status="covered", notes="Fixture task has language and evidence in backgroundb")
    for row in r["semantic_support"]:
        row.update(status="supported", notes="Fixture block exactly repeats source fact")
    r["pages"][0].update(status="pass", notes="test fixture")
    write_json(out / "audit/review.json", r)
    return out


def test_finalize_requires_review_identity(reviewable):
    r = read_json(reviewable / "audit/review.json")
    r["build_id"] = "wrong"
    write_json(reviewable / "audit/review.json", r)
    with pytest.raises(ValueError, match="identity"):
        finalize(reviewable)


def test_finalize_requires_every_page(reviewable):
    r = read_json(reviewable / "audit/review.json")
    r["pages"] = []
    write_json(reviewable / "audit/review.json", r)
    with pytest.raises(ValueError, match="Every rendered page"):
        finalize(reviewable)


def test_changed_docx_invalidates_review(reviewable):
    with (reviewable / "output.docx").open("ab") as f:
        f.write(b"changed")
    with pytest.raises(ValueError, match="Stale"):
        finalize(reviewable)


def test_changed_generator_invalidates_review(reviewable, monkeypatch):
    monkeypatch.setattr("word_skills.delivery.code_identity", lambda: "changed")
    with pytest.raises(ValueError, match="Generator changed"):
        finalize(reviewable)


def test_finalize_integrity_unit_fixture(reviewable):
    assert finalize(reviewable)["state"] == "reviewed"


def test_legacy_draft_warns_and_missing_case_is_structured(sample, tmp_path):
    assert "TASK_COVERAGE_MISSING" in codes(validate(sample[2], sample[1]))
    result = validate(sample[2], tmp_path / "empty")
    assert not result["passed"]
    assert {"NO_SOURCES", "NO_EVIDENCE"} <= codes(result)
    out = tmp_path / "blocked"
    assert build(sample[3], tmp_path / "empty", out, no_preview=True)["state"] == "blocked"
    assert not (out / "output.docx").exists()


@pytest.mark.parametrize("mutation", [
    lambda d: d.update(task_requirements=None),
    lambda d: d.update(task_requirements=[]),
    lambda d: d.update(task_requirements=[{"id": "T1", "text": "任务", "acceptance": "明确依据", "block_ids": ["absent"]}]),
    lambda d: d.update(task_requirements=[{"id": "T1", "text": "任务", "block_ids": ["backgroundb"]}]),
])
def test_task_contract_rejects_untraced_or_incomplete_items(sample, mutation):
    mutation(sample[2])
    assert not validate(sample[2], sample[1])["passed"]


@pytest.mark.parametrize("field,status", [("task_coverage", "gap"), ("task_coverage", "partial"), ("semantic_support", "unsupported")])
def test_unresolved_content_cannot_finalize(reviewable, field, status):
    report = read_json(reviewable / "audit/review.json")
    report[field][0].update(status=status, notes="Concrete unresolved observation")
    write_json(reviewable / "audit/review.json", report)
    with pytest.raises(ValueError, match=field):
        finalize(reviewable)


def test_review_does_not_auto_approve_undetected_semantics(sample, tmp_path):
    from word_skills.review import review_rows, validate_review
    sample[2]["task_requirements"] = [{"id": "T1", "text": "说明技术栈", "acceptance": "引用实际来源", "block_ids": ["backgroundb"]}]
    sample[2]["sections"][0]["blocks"][0]["text"] += "系统性能大幅提升。"
    # Public numeric checks cannot decide entailment. Review remains pending.
    assert validate(sample[2], sample[1])["passed"]
    report = review_rows(sample[2])
    assert all(r["status"] == "pending" for rows in report.values() for r in rows)
    with pytest.raises(ValueError, match="task_coverage"):
        validate_review(sample[2], report)
    report["task_coverage"][0].update(status="covered", notes="Technical stack identified")
    for row in report["semantic_support"]:
        row.update(status="supported", notes="Fixture exact fact")
    report["semantic_support"][0].update(status="unsupported", notes="Source contains language version, no performance result")
    with pytest.raises(ValueError, match="semantic_support"):
        validate_review(sample[2], report)


def test_duplicate_semantic_review_cannot_hide_unreviewed_blocks(reviewable):
    report = read_json(reviewable / "audit/review.json")
    report["semantic_support"][-1] = report["semantic_support"][0].copy()
    write_json(reviewable / "audit/review.json", report)
    with pytest.raises(ValueError, match="semantic_support"):
        finalize(reviewable)


def test_skill_changes_invalidate_review(reviewable, monkeypatch):
    monkeypatch.setattr("word_skills.delivery.skill_identity", lambda: "changed")
    with pytest.raises(ValueError, match="Skill instructions"):
        finalize(reviewable)
