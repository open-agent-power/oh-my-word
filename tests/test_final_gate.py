import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from oh_my_patent import parser
from oh_my_patent.cli import main
from oh_my_patent.gate import CHECKS, evidence_digest, final_issues, review_skeleton, sha256
from oh_my_patent.lint import lint
from oh_my_patent.report import LintReport
from oh_my_patent.render import DocxRenderer
from oh_my_patent.schema import parse_claims_block

DRAFT = '''---
title: 一种温度采集方法
type: invention
---
# 权利要求书
1. 一种温度采集方法，其特征在于，包括：获取传感器温度值；将所述温度值存储到存储器。
# 说明书
## 技术领域
本发明涉及传感器数据采集技术领域。
## 背景技术
现有温度采集方式需要人工抄录，易发生记录遗漏。
## 发明内容
### 要解决的技术问题
解决人工抄录造成数据遗漏的问题。
### 技术方案
获取传感器温度值；将所述温度值存储到存储器。
### 有益效果
通过自动存储温度值减少人工记录遗漏。
## 具体实施方式
采集器与数字温度传感器连接，读取传感器输出的温度数值，并将读数写入存储器，以供后续读取。存储失败时返回错误，保留原读数供重试。
# 摘要
本发明涉及一种温度采集方法，获取传感器温度值并将其存储到存储器，通过自动存储减少人工记录遗漏，用于传感器数据采集。
'''


class FinalGateTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.path = self.root / 'draft.md'
        self.path.write_text(DRAFT, encoding='utf8')
        self.draft = parser.load(self.path)
        self.source = self.root / 'technical.txt'
        self.source.write_text('获取传感器温度值。将所述温度值存储到存储器。', encoding='utf8')
        self.facts_path = self.root / 'TECH_FACTS.json'
        self.facts = {'schema_version': 1, 'sources': [{'id': 'S1', 'path': 'technical.txt', 'sha256': sha256(self.source)}],
                      'facts': [{'id': 'F1', 'text': '获取传感器温度值', 'status': 'verified', 'source_id': 'S1', 'locator': '第1句', 'quote': '获取传感器温度值'},
                                {'id': 'F2', 'text': '将所述温度值存储到存储器', 'status': 'verified', 'source_id': 'S1', 'locator': '第2句', 'quote': '将所述温度值存储到存储器'}]}
        self.save_facts()
        self.review_path = self.root / 'review.json'
        self.review = review_skeleton(self.path, self.facts_path, self.draft)
        self.review.update(reviewer='test fixture only', reviewed_at='2026-09-30')
        for check in self.review['checks']:
            check.update(status='pass', evidence='测试夹具中的结构化声明，不是实际专利认证')
        self.review['claim_support'] = [{'claim': 1, 'features': [
            {'text': '获取传感器温度值', 'description_quote': '获取传感器温度值', 'fact_ids': ['F1']},
            {'text': '将所述温度值存储到存储器', 'description_quote': '将所述温度值存储到存储器', 'fact_ids': ['F2']}]}]
        self.review['accepted_warnings'] = [{'code': i.code, 'reason': '测试夹具仅测试契约，不证明实际法律判断'} for i in lint(self.draft).warnings]
        self.save_review()

    def tearDown(self):
        self.tmp.cleanup()

    def save_facts(self):
        self.facts_path.write_text(json.dumps(self.facts, ensure_ascii=False), encoding='utf8')

    def save_review(self):
        self.review_path.write_text(json.dumps(self.review, ensure_ascii=False), encoding='utf8')

    def issues(self):
        return final_issues(self.path, self.draft, lint(self.draft), self.facts_path, self.review_path)

    def codes(self):
        return {x.code for x in self.issues()}

    def cli(self, *args):
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return main(list(args))

    def test_valid_contract(self):
        self.assertEqual(self.issues(), [])

    def test_missing_evidence(self):
        self.assertIn('FINAL-EVIDENCE', {x.code for x in final_issues(self.path, self.draft, LintReport())})

    def test_scaffold_pending(self):
        self.review = review_skeleton(self.path, self.facts_path, self.draft)
        self.save_review()
        self.assertIn('FINAL-CHECK', self.codes())

    def test_modified_draft(self):
        self.path.write_text(DRAFT + '\n变化', encoding='utf8')
        self.assertIn('FINAL-STALE', self.codes())

    def test_modified_source(self):
        self.source.write_text('已修改来源', encoding='utf8')
        self.assertIn('FINAL-STALE', self.codes())
        self.assertIn('FINAL-SOURCE', self.codes())

    def test_modified_facts(self):
        self.facts['facts'][0]['text'] += ' changed'
        self.save_facts()
        self.assertIn('FINAL-STALE', self.codes())

    def test_unverified_fact(self):
        self.facts['facts'][0]['status'] = 'partial'
        self.save_facts()
        self.assertIn('FINAL-SUPPORT', self.codes())

    def test_unknown_fact(self):
        self.review['claim_support'][0]['features'][0]['fact_ids'] = ['F999']
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())

    def test_forged_quote(self):
        self.facts['facts'][0]['quote'] = '不存在的阈值0.25'
        self.save_facts()
        self.assertIn('FINAL-FACT', self.codes())

    def test_wrong_description_quote(self):
        self.review['claim_support'][0]['features'][0]['description_quote'] = '不在说明书'
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())

    def test_missing_claim_support(self):
        self.review['claim_support'] = []
        self.save_review()
        self.assertIn('FINAL-CLAIM', self.codes())

    def test_duplicate_claim_support(self):
        self.review['claim_support'] *= 2
        self.save_review()
        self.assertIn('FINAL-CLAIM', self.codes())

    def test_empty_features(self):
        self.review['claim_support'][0]['features'] = []
        self.save_review()
        self.assertIn('FINAL-CLAIM', self.codes())

    def test_incomplete_check(self):
        self.review['checks'].pop()
        self.save_review()
        self.assertIn('FINAL-CHECK', self.codes())

    def test_na_needs_reason(self):
        self.review['checks'][0].update(status='not_applicable', evidence='')
        self.save_review()
        self.assertIn('FINAL-CHECK', self.codes())

    def test_duplicate_check(self):
        self.review['checks'].append(self.review['checks'][0])
        self.save_review()
        self.assertIn('FINAL-SCHEMA', self.codes())

    def test_p0_unresolved_or_accepted(self):
        for status in ['open', 'accepted']:
            self.review['findings'] = [{'severity': 'P0', 'status': status, 'evidence': '数量冲突尚未解决'}]
            self.save_review()
            self.assertIn('FINAL-FINDING', self.codes())

    def test_placeholder(self):
        self.path.write_text(DRAFT + '\nTODO', encoding='utf8')
        self.assertIn('FINAL-PLACEHOLDER', self.codes())

    def test_force_does_not_bypass_final(self):
        out = self.root / 'output.docx'
        code = self.cli('build', str(self.path), '-o', str(out), '--final', '--force')
        self.assertEqual(code, 2)
        self.assertFalse(out.exists())

    def test_build_and_handoff_cli(self):
        out = self.root / 'review.docx'
        args = ['--facts', str(self.facts_path), '--review', str(self.review_path)]
        self.assertEqual(self.cli('build', str(self.path), '-o', str(out), '--final', *args), 0)
        self.assertIn('温度采集', '\n'.join(p.text for p in Document(out).paragraphs))
        self.assertEqual(Document(out).styles['Normal'].font.size.pt, 12)
        self.assertEqual(Document(out).styles['Normal'].paragraph_format.line_spacing, 1.5)
        handoff = self.root / 'handoff'
        self.assertEqual(self.cli('handoff', str(self.path), '-o', str(handoff), *args), 0)
        manifest = json.loads((handoff / 'manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(manifest['status'], 'official_conversion_pending')
        self.assertEqual(manifest['filing_status'], 'not_submitted')
        for item in manifest['files']:
            self.assertEqual(item['sha256'], sha256(handoff / item['path']))
        # handoff 与 build 共用默认 auto：发明稿走固定上游，并在 manifest 记录实际制作器。
        self.assertEqual(manifest['renderer'], 'upstream')
        claims_docx = next((handoff / 'review-docx').glob('*权利要求书.docx'))
        self.assertEqual(Document(str(claims_docx)).paragraphs[0].style.name, 'Heading 1')
        local_handoff = self.root / 'handoff-local'
        self.assertEqual(
            self.cli('handoff', str(self.path), '-o', str(local_handoff), *args, '--renderer', 'local'),
            0,
        )
        local_manifest = json.loads((local_handoff / 'manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(local_manifest['renderer'], 'local')
        self.assertEqual(self.cli('handoff', str(self.path), '-o', str(handoff), *args), 1)

    def test_malformed_review_is_error(self):
        self.review['checks'] = None
        self.save_review()
        self.assertIn('FINAL-SCHEMA', self.codes())

    def test_empty_normalized_quote_is_rejected(self):
        from oh_my_patent.gate import _contains
        self.assertFalse(_contains('**', '无关正文'))
        for feature in self.review['claim_support'][0]['features']:
            feature.update(text='**', description_quote='**')
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())
        self.assertIn('FINAL-CLAIM', self.codes())

    def test_review_init_bad_source_reports_error(self):
        self.facts['sources'] = [{}]
        self.save_facts()
        self.assertEqual(self.cli('review-init', str(self.path), '--facts', str(self.facts_path), '-o', str(self.root/'new.json')), 1)

    def test_type_override_stales_review(self):
        from oh_my_patent.schema import PatentType
        self.draft.patent_type = PatentType.UTILITY
        self.assertIn('FINAL-STALE', self.codes())

    def test_extra_sections_support(self):
        from oh_my_patent.schema import Subsection
        self.draft.extra_sections.append(Subsection(heading='补充实施说明', body='独有的实施说明：获取传感器温度值'))
        self.review['claim_support'][0]['features'][0]['description_quote'] = '独有的实施说明：获取传感器温度值'
        self.save_review()
        self.assertNotIn('FINAL-SUPPORT', self.codes())

    def test_missing_one_feature_is_blocked(self):
        self.review['claim_support'][0]['features'].pop()
        self.save_review()
        self.assertIn('FINAL-COVERAGE', self.codes())

    def test_unrelated_verified_fact_is_blocked(self):
        self.review['claim_support'][0]['features'][0]['fact_ids'] = ['F2']
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())

    def test_unrelated_description_is_blocked(self):
        self.review['claim_support'][0]['features'][0]['description_quote'] = '解决人工抄录造成数据遗漏的问题。'
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())

    def test_paraphrase_requires_explicit_mapping(self):
        feature = self.review['claim_support'][0]['features'][0]
        feature['description_quote'] = '读取传感器输出的温度数值'
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())
        feature['support_basis'] = '获取温度值与读取输出数值在该测试材料中对应同一采集步骤，人工核对。'
        feature['fact_links'] = [{'fact_id': 'F1', 'source_quote': '获取传感器温度值',
                                 'basis': 'F1明确记载获取温度值，说明书展开读取操作。'}]
        self.save_review()
        self.assertNotIn('FINAL-SUPPORT', self.codes())
        feature['fact_links'][0]['source_quote'] = '不存在的采集步骤'
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())

    def test_mapping_must_cover_all_fact_ids(self):
        feature = self.review['claim_support'][0]['features'][0]
        feature.update(fact_ids=['F1', 'F2'], support_basis='测试夹具',
                       fact_links=[{'fact_id': 'F1', 'source_quote': '获取传感器温度值', 'basis': '测试夹具'}])
        self.save_review()
        self.assertIn('FINAL-SUPPORT', self.codes())

    def test_missing_preamble_feature_is_blocked(self):
        self.draft.claims[0].text = self.draft.claims[0].text.replace('，其特征在于', '，包括校准传感器，其特征在于')
        self.assertIn('FINAL-COVERAGE', self.codes())

    def test_technical_subject_not_silently_excluded(self):
        self.draft.claims[0].text = self.draft.claims[0].text.replace('一种温度采集方法', '一种采用激光测温的采集方法')
        self.assertIn('FINAL-COVERAGE', self.codes())

    def test_quote_including_intro_still_counts_as_coverage(self):
        feature = self.review['claim_support'][0]['features'][0]
        feature.update(text='其特征在于，包括：获取传感器温度值', support_basis='固定引导语不增加技术内容。',
                       fact_links=[{'fact_id': 'F1', 'source_quote': '获取传感器温度值', 'basis': '本条对应实际采集步骤。'}])
        self.save_review()
        self.assertEqual(self.issues(), [])

    def test_invalid_render_parameters_do_not_write_output(self):
        evidence = ['--facts', str(self.facts_path), '--review', str(self.review_path)]
        for flag, value in [('--font-size', '6'), ('--font-size', '0'), ('--font-size', '20'),
                            ('--font-size', 'nan'), ('--font-size', 'inf'),
                            ('--line-spacing', '0'), ('--line-spacing', '0.2'),
                            ('--line-spacing', '-1'), ('--line-spacing', 'nan')]:
            for command, mode in [('build', []), ('build', ['--final']), ('build', ['--force']), ('handoff', [])]:
                with self.subTest(flag=flag, value=value, command=command, mode=mode):
                    out = self.root / 'bad-output'
                    self.assertNotEqual(self.cli(command, str(self.path), '-o', str(out), flag, value, *mode, *evidence), 0)
                    self.assertFalse(out.exists())

    def test_coverage_failure_blocks_build_and_handoff(self):
        self.review['claim_support'][0]['features'].pop()
        self.save_review()
        for command, mode in [('build', ['--final']), ('handoff', [])]:
            out = self.root / 'missing-support-output'
            self.assertEqual(self.cli(command, str(self.path), '-o', str(out), *mode,
                                      '--facts', str(self.facts_path), '--review', str(self.review_path)), 2)
            self.assertFalse(out.exists())

    def test_raw_latex_is_blocked(self):
        self.path.write_text(DRAFT + '\n$$ x \\frac{a}{b} $$', encoding='utf8')
        self.assertIn('FINAL-MATH', self.codes())

    def test_table_native_render(self):
        self.draft.embodiments += '\n| 项目 | 数值 |\n| --- | --- |\n| 温度 | 25 |'
        doc = DocxRenderer(self.draft).render()
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(1,1).text, '25')
        self.assertFalse(any('| --- |' in p.text for p in doc.paragraphs))

    def test_bad_table_columns_rejected(self):
        self.draft.embodiments += '\n| A | B |\n| --- | --- |\n| only one |'
        with self.assertRaises(ValueError):
            DocxRenderer(self.draft).render()

    def test_final_corrupt_image(self):
        from oh_my_patent.schema import Drawing
        image = self.root/'bad.png';image.write_text('not an image')
        self.draft.drawings = [Drawing(number=1, path=str(image), is_abstract=True)]
        self.assertIn('FINAL-DRAW', self.codes())

    def test_figure_mutation_stales_review(self):
        from oh_my_patent.schema import Drawing
        from PIL import Image
        image = self.root/'figure.png';Image.new('RGB',(20,20)).save(image)
        self.draft.drawings = [Drawing(number=1, path=str(image), is_abstract=True)]
        self.review['evidence_sha256'] = evidence_digest(self.facts_path, self.draft)
        self.save_review()
        Image.new('RGB',(30,30)).save(image)
        self.assertIn('FINAL-STALE', self.codes())

    def test_abstract_figure_split(self):
        from oh_my_patent.schema import Drawing
        from PIL import Image
        image = self.root/'figure.png';Image.new('RGB',(20,20)).save(image)
        self.draft.drawings = [Drawing(number=1, path=str(image), is_abstract=True)]
        parts = dict(DocxRenderer(self.draft).render_split())
        ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        self.assertEqual(len(parts['abstract'].element.findall('.//w:drawing',ns)), 0)
        self.assertEqual(len(parts['abstract_drawing'].element.findall('.//w:drawing',ns)), 1)


class PreservationTest(unittest.TestCase):
    def test_gap_not_repaired(self):
        claims = parse_claims_block('1. 一种方法，其特征在于，采集数据。\n3. 根据权利要求1所述的方法，其特征在于，存储数据。')
        self.assertEqual([c.number for c in claims], [1, 3])
        draft = parser.parse_markdown(DRAFT.replace('1. 一种温度采集方法', '3. 一种温度采集方法'))
        self.assertIn('CLAIM-002', {x.code for x in lint(draft).issues})

    def test_missing_figure_visible_in_draft_and_final_blocked(self):
        draft = parser.parse_markdown(DRAFT.replace('type: invention', 'type: invention\ndrawings:\n  - {number: 1, caption: 图1为结构图。}'))
        text = '\n'.join(p.text for p in DocxRenderer(draft).render().paragraphs)
        self.assertIn('图1缺失', text)
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / 'draft.md';p.write_text(DRAFT, encoding='utf8')
            codes = {x.code for x in final_issues(p, draft, lint(draft))}
            self.assertIn('FINAL-DRAW', codes)
