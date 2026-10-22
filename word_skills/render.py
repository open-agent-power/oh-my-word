"""One renderer, two content modes, and explicit bounded template support."""
from __future__ import annotations

import re
from defusedxml import ElementTree
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .common import check_docx, inside, sha256
from .inline import inline_runs
from .models import validate_renderable_draft

STYLE_NAMES = ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]
SECTION_ATTRS = ["page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin", "header_distance", "footer_distance"]


def fingerprint(path):
    check_docx(path)
    doc = Document(path)
    warnings = []
    with ZipFile(path) as z:
        xml = ElementTree.fromstring(z.read("word/document.xml"))
        tags = {element.tag.split("}")[-1] for element in xml.iter()}
        for marker, label in [("fldChar", "Fields"), ("numPr", "Numbering"), ("txbxContent", "Text boxes"), ("ins", "Tracked changes"), ("del", "Tracked changes"), ("sdt", "Content controls"), ("drawing", "Drawings")]:
            if marker in tags:
                warnings.append(f"{label}: manual fidelity review required")
    if len(doc.sections) != 1:
        warnings.append("Multiple sections: not supported by style-only template import")
    styles = {}
    for name in STYLE_NAMES:
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        fonts = style.element.find(".//" + qn("w:rFonts"))
        styles[name] = {"font": style.font.name, "east_asia": fonts.get(qn("w:eastAsia")) if fonts is not None else None,
                        "size_pt": style.font.size.pt if style.font.size else None,
                        "bold": style.font.bold, "outline_level": name[-1] if name.startswith("Heading") else None}
    return {"schema_version": 1, "template_sha256": sha256(path),
            "sections": [{a: int(getattr(s, a)) if getattr(s, a) is not None else None for a in SECTION_ATTRS} for s in doc.sections],
            "styles": styles, "table_styles": sorted({t.style.name for t in doc.tables if t.style}),
            "headers": [[p.text for p in s.header.paragraphs] for s in doc.sections],
            "footers": [[p.text for p in s.footer.paragraphs] for s in doc.sections],
            "warnings": warnings,
            "limits": "Basic section/paragraph style fingerprint; not a full template reconstruction or pixel fidelity claim"}


def font(style, name, size, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        if "Theme" in attr:
            del rfonts.attrib[attr]
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:hint"), "eastAsia")


def default_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(23)
    font(doc.styles["Normal"], "宋体", 11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.35
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    font(doc.styles["Title"], "黑体", 20, True)
    # Some python-docx default templates carry a blue bottom border on Title.
    # New documents use a plain title; imported customer templates stay intact.
    for border in list(doc.styles["Title"].element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.styles["Title"].font.underline = False
    for level, size in [(1, 15), (2, 12), (3, 11)]:
        style = doc.styles[f"Heading {level}"]
        font(style, "黑体", size, True)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def clear_body(doc):
    for item in list(doc.element.body):
        if item.tag != qn("w:sectPr"):
            doc.element.body.remove(item)


def append_content(doc, draft, title=True, asset_root=None):
    if draft.get("mode") in {"copyright-source", "copyright-manual"} and draft.get("software_name"):
        for s in doc.sections:
            header = s.header.paragraphs[0]
            header.clear()
            header.style = doc.styles["Normal"]
            header.paragraph_format.tab_stops.clear_all()
            header.paragraph_format.tab_stops.add_tab_stop(
                s.page_width - s.left_margin - s.right_margin, WD_TAB_ALIGNMENT.RIGHT)
            run = header.add_run(f"{draft['software_name']} {draft['software_version']}\t第 ")
            run.font.size = Pt(9)
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            header._p.append(field)
            header.add_run(" 页").font.size = Pt(9)
            for footer in s.footer.paragraphs:
                footer.clear()
    if title:
        doc.add_paragraph(draft["title"], "Title")
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for section in draft["sections"]:
        doc.add_paragraph(section["title"], f"Heading {section.get('level', 1)}")
        for block in section["blocks"]:
            if block["type"] == "code":
                # Preserve every source line, including blank lines and indentation.
                # Wrapped visual lines are measured by the PDF, never estimated here.
                for line in block["text"].splitlines():
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1
                    p.paragraph_format.widow_control = False
                    run = p.add_run(line)
                    run.font.name = "Consolas"
                    run.font.size = Pt(8)
            elif block["type"] == "paragraph":
                level = block.get("heading_level")
                p = doc.add_paragraph(style=f"Heading {level}" if level in (1, 2, 3) else None)
                if block.get("inline_format") == "markdown":
                    for text, is_code in inline_runs(block["text"]):
                        run = p.add_run(text)
                        if is_code:
                            run.font.name = "Consolas"
                            rpr = run._element.get_or_add_rPr()
                            rfonts = rpr.rFonts
                            if rfonts is None:
                                rfonts = OxmlElement("w:rFonts")
                                rpr.insert(0, rfonts)
                            rfonts.set(qn("w:eastAsia"), "Consolas")
                else:
                    p.add_run(block["text"])
                p.paragraph_format.widow_control = True
                if block.get("monospace"):
                    p.paragraph_format.line_spacing = 1
                    for run in p.runs:
                        run.font.name, run.font.size = "Consolas", Pt(8)
            elif block["type"] == "image":
                image = inside(asset_root, block["path"])
                if sha256(image) != block["sha256"]:
                    raise ValueError("Image differs from the audited source")
                from docx.image.image import Image
                info = Image.from_file(str(image))
                available_height = doc.sections[0].page_height - doc.sections[0].top_margin - doc.sections[0].bottom_margin - Mm(20)
                picture_width = min(width, int(available_height * info.px_width / info.px_height))
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(image), width=picture_width)
                doc.add_paragraph(block["text"])
            else:
                headers = block["headers"]
                table = doc.add_table(rows=1, cols=len(headers))
                chosen = block.get("style", "Table Grid")
                if chosen not in doc.styles or doc.styles[chosen].type != 3:
                    raise ValueError(f"Unknown table style: {chosen}")
                table.style = chosen
                table.autofit = False
                weights = block.get("widths", [1] * len(headers))
                total = sum(weights)
                widths = [int(width * w / total) for w in weights]
                for col, w in zip(table.columns, widths):
                    col.width = w
                for cell, text, w in zip(table.rows[0].cells, headers, widths):
                    cell.text, cell.width = str(text), w
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row in block["rows"]:
                    for cell, text, w in zip(table.add_row().cells, row, widths):
                        cell.text, cell.width = str(text), w
                # Keep ordinary rows intact so labels and their values do not
                # get separated across pages. A row taller than a page can
                # still split in Word/LibreOffice despite this OOXML hint.
                for row in table.rows:
                    cant_split = OxmlElement("w:cantSplit")
                    row._tr.get_or_add_trPr().append(cant_split)
                borders = OxmlElement("w:tblBorders")
                for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    border = OxmlElement("w:" + side)
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "4")
                    border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                for r, row in enumerate(table.rows):
                    for cell in row.cells:
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        props = cell._tc.get_or_add_tcPr()
                        margins = OxmlElement("w:tcMar")
                        for side in ("top", "bottom", "left", "right"):
                            n = OxmlElement("w:" + side)
                            n.set(qn("w:w"), "90")
                            n.set(qn("w:type"), "dxa")
                            margins.append(n)
                        props.append(margins)
                        if r == 0:
                            fill = OxmlElement("w:shd")
                            fill.set(qn("w:fill"), "E8EDF2")
                            props.append(fill)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(3)
                            paragraph.paragraph_format.space_before = Pt(3)
                            for run in paragraph.runs:
                                run.bold = r == 0


def render(draft, output, template=None, template_mode="style", *, asset_root=None):
    draft = validate_renderable_draft(draft)
    profile = None
    if template:
        profile = fingerprint(template)
        if len(profile["sections"]) != 1:
            raise ValueError("V1 template support is single-section; multiple sections require a dedicated adapter")
    if template and template_mode == "slots":
        from docxtpl import DocxTemplate
        from jinja2 import StrictUndefined
        from jinja2.sandbox import SandboxedEnvironment
        with ZipFile(template) as archive:
            xml = "\n".join(archive.read(n).decode("utf-8") for n in archive.namelist() if n.startswith("word/") and n.endswith(".xml"))
        tags = re.findall(r"\{\{.*?\}\}|\{%.*?%\}", xml, re.S)
        allowed = {"{{ title }}", "{{p body }}"}
        template_doc = Document(template)
        paragraphs = [p.text for p in template_doc.paragraphs]
        if (any(t not in allowed for t in tags) or tags.count("{{p body }}") != 1
                or tags.count("{{ title }}") != 1 or paragraphs.count("{{p body }}") != 1
                or paragraphs.count("{{ title }}") != 1):
            raise ValueError("Slot template requires exactly one standalone {{ title }} and one standalone {{p body }} paragraph")
        tpl = DocxTemplate(template)
        body = tpl.new_subdoc()
        for a in SECTION_ATTRS:
            setattr(body.sections[0], a, getattr(template_doc.sections[0], a))
        append_content(body, draft, title=False, asset_root=asset_root)
        tpl.render({"title": draft["title"], "body": body}, jinja_env=SandboxedEnvironment(undefined=StrictUndefined), autoescape=True)
        tpl.save(output)
    else:
        doc = Document(template) if template else default_document()
        if template:
            clear_body(doc)  # Explicit --template-mode style: reference body is not an editable input document.
        append_content(doc, draft, asset_root=asset_root)
        doc.core_properties.title = draft["title"]
        doc.core_properties.author = ""
        doc.save(output)
    return profile
