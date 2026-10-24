"""Exercise an installed wheel outside the checkout using synthetic materials.

Run with the clean environment's Python and -I. This does not review content,
approve pages, call a model, or contact a knowledge service.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
from pathlib import Path
import subprocess
import sys


ROOT = Path(__file__).resolve().parents[1]


def require_outside(path, root=ROOT):
    path = Path(path).resolve()
    if path.is_relative_to(root.resolve()):
        raise ValueError("Installed-wheel checks must run outside the source checkout")
    return path


def command(args, cwd, records, expected=0):
    result = subprocess.run(args, cwd=cwd, capture_output=True, text=True,
                            encoding="utf-8", errors="replace", timeout=180)
    records.append({"command": [str(a) for a in args], "exit_code": result.returncode,
                    "stdout": result.stdout, "stderr": result.stderr})
    if result.returncode != expected:
        raise RuntimeError(f"Unexpected exit {result.returncode}, expected {expected}: {args}")
    return result.stdout


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


#: 专利制作器安装态夹具。纯文本、无图片，避免依赖外部图件或本机 Office。
PATENT_RENDERER_FIXTURE = """---
title: 安装验收专利样稿
type: invention
---

# 权利要求书

1. 一种安装验收方法，其特征在于，包括执行一次固定上游制作。

# 说明书

【技术领域】

本材料仅用于安装验收，标记INSTALL_FIELD。

【背景技术】

不含真实技术事实，标记INSTALL_BG。

【发明内容】

包括接收步骤与输出步骤。

【具体实施方式】

步骤S101，接收记录，标记INSTALL_END。

# 摘要

本材料仅用于安装验收，标记INSTALL_ABS。
"""


def check_patent_renderers(output, records, scripts, suffix, Document):
    """在安装态下分别执行默认本地制作器与固定上游制作器。

    安装态最容易漏掉的是 ``_vendor`` 打包与上游脚本的导入路径；这里直接
    调用安装后的 ``oh-my-patent`` 控制台脚本，并核对固定原件哈希、A4 版面和
    三类文书，避免只验证 import 而不验证真实制作。
    """
    from oh_my_patent.render.upstream import upstream_identity

    patent_cli = str(scripts / ("oh-my-patent" + suffix))
    fixture = output / "patent-renderer-fixture.md"
    fixture.write_text(PATENT_RENDERER_FIXTURE, encoding="utf-8")

    upstream_docx = output / "patent-upstream.docx"
    upstream_stdout = command(
        [patent_cli, "build", str(fixture), "-o", str(upstream_docx), "--renderer", "upstream"],
        output,
        records,
    )
    if "上游制作器" not in upstream_stdout:
        raise RuntimeError("Upstream renderer did not report its pinned identity")
    identity = upstream_identity()
    if identity["commit"] != "6935a14d342ed9c7a07440c999d1b7bf1852a701":
        raise RuntimeError(f"Unexpected pinned upstream commit: {identity['commit']}")
    if identity.get("modifications"):
        raise RuntimeError("Vendored upstream files are marked as modified")

    document = Document(upstream_docx)
    section = document.sections[0]
    page_mm = [round(section.page_width.mm, 1), round(section.page_height.mm, 1)]
    if abs(page_mm[0] - 210) > 1 or abs(page_mm[1] - 297) > 1:
        raise RuntimeError(f"Upstream renderer did not apply A4: {page_mm}")
    titles = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    if titles != ["权利要求书", "说明书", "说明书摘要"]:
        raise RuntimeError(f"Unexpected upstream part titles: {titles}")
    body = "\n".join(p.text for p in document.paragraphs)
    missing = [m for m in ("INSTALL_FIELD", "INSTALL_BG", "INSTALL_END", "INSTALL_ABS") if m not in body]
    if missing:
        raise RuntimeError(f"Upstream renderer lost fixture markers: {missing}")
    if "$" in body:
        raise RuntimeError("Upstream renderer left raw dollar math in the body")

    # 默认路径必须仍是本地制作器：不带 --renderer 也应成功，且不改变既有行为。
    local_docx = output / "patent-local.docx"
    command([patent_cli, "build", str(fixture), "-o", str(local_docx)], output, records)
    if not local_docx.is_file():
        raise RuntimeError("Default (local) patent renderer produced no file")

    return {
        "upstream_commit": identity["commit"],
        "upstream_files": len(identity["files"]),
        "upstream_titles": titles,
        "upstream_page_mm": page_mm,
        "default_renderer": "local",
        "scope": "synthetic fixtures; no writing-quality or filing acceptance",
    }



def smoke(output, records, *, preview=False, soffice=None, pdftoppm=None):
    import word_skills
    import oh_my_patent
    from docx import Document
    from word_skills.evidence import contract

    modules = {m.__name__: str(require_outside(m.__file__)) for m in (word_skills, oh_my_patent)}
    cli = [sys.executable, "-I", "-X", "utf8", "-m", "word_skills"]
    command([sys.executable, "-I", "-m", "pip", "check"], output, records)
    command([*cli, "skills", "list", "--json"], output, records)
    command([*cli, "route", "请写技术报告", "--json"], output, records)
    command([*cli, "patent", "template", "invention", "-o", str(output / "patent-draft.md")], output, records)
    scripts = Path(sys.executable).parent
    # Exercise the installed console scripts as well as the module dispatcher.
    suffix = ".exe" if sys.platform == "win32" else ""
    for name in ("word-skills", "oh-my-patent"):
        command([str(scripts / (name + suffix)), "--help"], output, records)
    patent_renderers = check_patent_renderers(output, records, scripts, suffix, Document)

    text = "这是安装验收使用的合成材料，不代表真实项目或专业审阅。"
    source = output / "source.txt"
    source.write_text(text + "\n", encoding="utf-8")
    case = output / "case"
    command([*cli, "ingest", str(source), "--case", str(case)], output, records)
    write_json(case / "evidence.json", {"schema_version": 1, "claims": [{
        "id": "C1", "statement": text, "source_id": "S001", "locators": ["L1"],
        "quote": text, "status": "verified", "verified_by": "synthetic-install-fixture",
        "verified_at": datetime.now(timezone.utc).isoformat()}]})
    sections = [{"id": sid, "title": title, "blocks": [{"id": sid + "-text",
                "type": "paragraph", "intent": "fact", "text": text, "claim_ids": ["C1"]}]}
                for sid, title in contract("technical-report")["required_sections"].items()]
    draft = {"schema_version": 1, "mode": "technical-report", "title": "安装验收样稿",
             "sections": sections, "task_requirements": [{"id": "T1", "text": "生成合成材料样稿",
             "acceptance": "保留合成材料标记和全部章节", "block_ids": [s["blocks"][0]["id"] for s in sections]}]}
    draft_path = output / "draft.json"
    write_json(draft_path, draft)
    command([*cli, "lint", str(draft_path), "--case", str(case)], output, records)
    delivery = output / "delivery"
    build = [*cli, "build", str(draft_path), "--case", str(case), "--out", str(delivery)]
    if preview:
        if soffice:
            build += ["--soffice", soffice]
        if pdftoppm:
            build += ["--pdftoppm", pdftoppm]
    else:
        build += ["--no-preview"]
    command(build, output, records)
    manifest = json.loads((delivery / "manifest.json").read_text("utf-8"))
    expected_state = "awaiting-review" if preview else "draft"
    if manifest["state"] != expected_state:
        raise RuntimeError(f"Expected {expected_state}, got {manifest['state']}")
    if text not in [p.text for p in Document(delivery / "output.docx").paragraphs]:
        raise RuntimeError("Installed build lost fixture text")
    # Never manufacture review notes to make the installation test go green.
    command([*cli, "finalize", str(delivery)], output, records, expected=2)
    if json.loads((delivery / "manifest.json").read_text("utf-8"))["state"] != expected_state:
        raise RuntimeError("Unreviewed fixture was promoted")
    if preview:
        pages = manifest["renderer"]["page_images"]
        if not pages or any(not (delivery / p).is_file() for p in pages):
            raise RuntimeError("Preview page images missing")

    project = output / "synthetic-project"
    project.mkdir()
    (project / "main.py").write_text("def greet():\n    return 'synthetic install fixture'\n", encoding="utf-8")
    command([*cli, "copyright", "prepare", str(project), "--out", str(output / "copyright-case"),
             "--name", "安装验收示例", "--version", "0.0.0"], output, records)
    return {"modules": modules, "version": importlib.metadata.version("word-skills"),
            "patent_renderers": patent_renderers,
            "delivery_state": manifest["state"], "pages": manifest.get("renderer", {}).get("pages"),
            "preview_executed": preview, "semantic_review": "not performed",
            "visual_review": "not performed", "fixture": "synthetic, not writing-quality evidence",
            "skill_bundle": "not included in wheel; use full repository for Agent Skill resources"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True, help="Fresh directory outside the checkout")
    parser.add_argument("--preview", action="store_true", help="Require real LibreOffice/Poppler conversion")
    parser.add_argument("--soffice")
    parser.add_argument("--pdftoppm")
    args = parser.parse_args()
    output = require_outside(args.out)
    output.mkdir(parents=True, exist_ok=False)
    report = {"status": "fail", "python": sys.version, "interpreter": sys.executable,
              "script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
              "resolved_dependencies": dict(sorted((d.metadata["Name"], d.version)
                                                   for d in importlib.metadata.distributions())),
              "commands": []}
    try:
        report.update(smoke(output, report["commands"], preview=args.preview,
                            soffice=args.soffice, pdftoppm=args.pdftoppm))
        report["status"] = "pass"
    except Exception as exc:
        report["error"] = f"{type(exc).__name__}: {exc}"
    write_json(output / "installed-delivery.json", report)
    print(json.dumps({"status": report["status"], "report": str(output / "installed-delivery.json"),
                      "error": report.get("error")}, ensure_ascii=False))
    return 0 if report["status"] == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
