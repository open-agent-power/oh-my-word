from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

import pdfplumber
from docx import Document
from pypdf import PdfReader

from .common import check_docx, finding, sha256
from .evidence import block_text


def inspect_docx(path, draft):
    check_docx(path)
    doc = Document(path)
    issues = []
    texts = [p.text for p in doc.paragraphs]
    texts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    texts += [p.text for s in doc.sections for container in (s.header, s.footer) for p in container.paragraphs]
    text = "\n".join(texts)
    if not text.strip():
        issues.append(finding("DOCX_EMPTY", "Document text empty"))
    if draft["title"] not in texts:
        issues.append(finding("DOCX_TITLE", "Document title missing"))
    code_lines = {line for section in draft["sections"] for block in section["blocks"] if block["type"] == "code" for line in block["text"].splitlines()}
    prose = "\n".join(t for t in texts if t not in code_lines)
    if re.search(r"TODO|XXX|\{\{|\{%|待补充", prose):
        issues.append(finding("DOCX_PLACEHOLDER", "Unresolved placeholder in final document"))
    for section in draft["sections"]:
        if section["title"] not in texts:
            issues.append(finding("DOCX_SECTION", f"Heading missing: {section['title']}"))
        for block in section["blocks"]:
            if block["type"] == "code":
                required = block["text"].splitlines()
            elif block["type"] in {"paragraph", "image"}:
                required = [block["text"]]
            else:
                required = [str(c) for row in [block["headers"]] + block["rows"] for c in row]
            if any(t not in texts for t in required):
                issues.append(finding("DOCX_CONTENT", f"Content missing: {block['id']}"))
    available = min(s.page_width - s.left_margin - s.right_margin for s in doc.sections)
    for i, table in enumerate(doc.tables, 1):
        if sum(col.width or 0 for col in table.columns) > available + 12700:
            issues.append(finding("TABLE_OVERFLOW", f"Table {i} grid wider than page content"))
    return issues


def preview(docx, folder, soffice=None, pdftoppm=None):
    """Run only supplied/local executables; no download or cloud fallback."""
    soffice = soffice or shutil.which("soffice") or shutil.which("soffice.exe")
    pdftoppm = pdftoppm or shutil.which("pdftoppm")
    if not soffice:
        raise RuntimeError("LibreOffice unavailable: pass --soffice PATH")
    if not pdftoppm:
        raise RuntimeError("Poppler unavailable: pass --pdftoppm PATH")
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    if (folder / "preview.pdf").exists() or any(folder.glob("page-*.png")):
        raise ValueError("Preview artifacts already exist; use a fresh output directory")
    with tempfile.TemporaryDirectory(prefix="word-skills-lo-") as temp:
        work = Path(temp)
        profile = (work / "profile").as_uri()
        command = [str(soffice), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(Path(docx).resolve())]
        result = subprocess.run(command, capture_output=True, timeout=90, text=True, encoding="utf-8", errors="replace")
        pdf = work / (Path(docx).stem + ".pdf")
        if result.returncode or not pdf.is_file() or pdf.stat().st_size == 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stdout} {result.stderr}")
        destination = folder / "preview.pdf"
        raster_command = [str(pdftoppm), "-png", "-r", "110", str(pdf), str(work / "page")]
        raster = subprocess.run(raster_command, capture_output=True, timeout=90)
        images = sorted(work.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
        count = len(PdfReader(pdf).pages)
        if raster.returncode or not count or [p.name for p in images] != [f"page-{i:0{len(str(count))}d}.png" for i in range(1, count + 1)]:
            raise RuntimeError("PDF page image rendering failed or incomplete")
        shutil.copyfile(pdf, destination)
        for image in images:
            shutil.copyfile(image, folder / image.name)
        page_names = [p.name for p in images]
    try:
        version = subprocess.run([str(soffice), "--version"], capture_output=True, text=True, timeout=20).stdout.strip() or None
    except (OSError, subprocess.TimeoutExpired):
        version = None
    return {"backend": "LibreOffice", "version": version, "pages": count, "page_images": page_names,
            "provenance": {"docx_sha256": sha256(docx), "pdf_sha256": sha256(destination),
                           "conversion_command": command, "raster_command": raster_command, "dpi": 110,
                           "tools": {name: {"path": str(Path(path).resolve()), "sha256": sha256(path)}
                                     for name, path in (("soffice", soffice), ("pdftoppm", pdftoppm))},
                           "page_hashes": {name: sha256(folder / name) for name in page_names}}}


def inspect_pdf(path, draft):
    issues, pages, fonts, extracted = [], [], set(), []
    normalize = lambda text: re.sub(r"\s+", "", str(text or ""))
    table_rows = Counter()
    headings = {s["title"].replace(" ", "") for s in draft["sections"]}
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            issues.append(finding("PDF_EMPTY", "No PDF pages"))
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            extracted.append(text)
            for table in page.extract_tables():
                table_rows.update(tuple(normalize(cell) for cell in row) for row in table)
            pages.append({"page": i, "characters": len(text), "width": page.width, "height": page.height})
            if len(re.sub(r"[\s\d]", "", text)) < 3:
                issues.append(finding("BLANK_PAGE", f"Page {i} blank or nearly empty"))
            if "\ufffd" in text or "(cid:" in text:
                issues.append(finding("PDF_GLYPH", f"Page {i} contains replacement glyph markers"))
            lines = [s.strip().replace(" ", "") for s in text.splitlines() if s.strip() and not s.strip().isdigit()]
            if lines and lines[-1] in headings:
                issues.append(finding("ORPHAN_HEADING", f"Heading at bottom of page {i}"))
            for char in page.chars:
                fonts.add(char.get("fontname", "unknown"))
                if char["x0"] < -1 or char["x1"] > page.width + 1 or char["top"] < -1 or char["bottom"] > page.height + 1:
                    issues.append(finding("PDF_BOUNDS", f"Text outside page {i}"))
                    break
    full_text = normalize("\n".join(extracted))
    expected = [("title", draft["title"])]
    required_rows = Counter()
    for section in draft["sections"]:
        expected.append((section["id"], section["title"]))
        for block in section["blocks"]:
            if block["type"] == "table":
                # Spatial cells preserve column order when visual text lines interleave.
                # Accumulate counts across blocks: one row cannot satisfy many copies.
                rows = Counter(tuple(normalize(cell) for cell in row) for row in block["rows"])
                required_rows.update(rows)
                header = tuple(normalize(cell) for cell in block["headers"])
                missing = sum(max(0, required_rows[row] - table_rows[row]) for row in rows)
                if missing or not table_rows[header]:
                    issues.append(finding("PDF_CONTENT_REVIEW",
                        f"Spatial extraction could not confirm table header or {missing} required row occurrence(s); check all cells and row counts, including page splits",
                        "warning", location=block["id"]))
            else:
                if any(normalize(line) not in full_text for line in block["text"].splitlines() if line.strip()):
                    issues.append(finding("PDF_CONTENT_REVIEW",
                        "Body text not confirmed by extraction; inspect for missing content or interleaving/page splits",
                        "warning", location=block["id"]))
    for location, text in expected:
        if normalize(text) not in full_text:
            issues.append(finding("PDF_CONTENT", "Expected content not found in extracted PDF; inspect rendering and extraction", location=location))
    return {"issues": issues, "pages": pages, "rendered_fonts": sorted(fonts),
            "font_substitution": "not automatically determined; compare rendered fonts and page images with template",
            "limits": "Text and spatial table extraction are heuristics, not completeness proof. Unconfirmed body/table content requires explicit warning review; page splits, overlaps, cell clipping and meaning require visual/content review"}
