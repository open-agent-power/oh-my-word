"""上游专利 Word 转换模块复用的验收测试。

断言分四类，避免把「生成成功」当成「可交付」：

1. 固定身份：vendored 原件与 ``UPSTREAM.json``、第三方许可逐字节一致；
2. 投影：稿件 → 上游四类 Markdown 的章节顺序和内容；
3. 边界版面：A4、黑色标题、宋体、字号、行距、首行缩进、纯文本编号；
4. 失败门禁：坏公式回退、缺图、普通括号 LaTeX、外观设计均明确失败。
"""

from __future__ import annotations

import base64
import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from oh_my_patent.cli import main as cli_main
from oh_my_patent.render import (
    UpstreamRenderError,
    markdown_parts,
    render_docx_upstream,
    upstream_identity,
)
from oh_my_patent.schema import Claim, Drawing, PatentDraft, PatentType
from oh_my_patent.spec import HEADING_NAME_FONT_SIZE_PT, RenderOptions, TITLE_FONT_SIZE_PT

REPO_ROOT = Path(__file__).resolve().parent.parent
VENDOR = REPO_ROOT / "oh_my_patent" / "render" / "_vendor" / "patent_application"
M_OMML = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"

#: 最小合法 PNG（1×1），避免测试依赖 Pillow。
PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _image(root: Path) -> Path:
    path = root / "figure-1.png"
    path.write_bytes(PNG_BYTES)
    return path


def _draft(root: Path, *, math: bool = False, drawings: bool = False) -> PatentDraft:
    embodiments = "步骤S101，接收带序号的记录。\n\n步骤S102，写入缓冲队列，标记END_D。"
    if math:
        embodiments += (
            "\n\n内联均值为 $\\bar{x}=\\frac{1}{n}\\sum_{i=1}^{n}x_i$，单位为毫秒。"
            "\n\n$$\nE=\\sum_{i=1}^{n}(x_i-\\mu)^2\n$$"
        )
    drawing_list = []
    if drawings:
        drawing_list = [Drawing(number=1, caption="队列处理结构图", path=str(_image(root)))]
    return PatentDraft(
        title="一种缓冲队列处理方法",
        patent_type=PatentType.INVENTION,
        technical_field="本发明涉及数据处理领域，标记FIELD_A。",
        background="输入记录经队列暂存后按既定顺序读取，标记BACKGROUND_B。",
        problems="读取顺序不稳定。",
        solution="接收模块读取记录，缓冲模块保存记录，输出模块返回记录。",
        effects="读取顺序可复现。",
        embodiments=embodiments,
        claims=[
            Claim(
                number=1,
                text="一种缓冲队列处理方法，其特征在于，包括接收记录、保存记录和按序号输出记录。",
            )
        ],
        abstract="本合成示例包括接收模块、缓冲模块和输出模块，标记ABSTRACT_E。",
        drawings=drawing_list,
    )


class IdentityTests(unittest.TestCase):
    def test_vendor_bytes_match_manifest(self) -> None:
        manifest = upstream_identity(verify=True)
        self.assertEqual(manifest["commit"], "6935a14d342ed9c7a07440c999d1b7bf1852a701")
        self.assertEqual(manifest["license"], "MIT")
        self.assertEqual(manifest["modifications"], [])
        for name, expected in manifest["files"].items():
            actual = hashlib.sha256((VENDOR / name).read_bytes()).hexdigest()
            self.assertEqual(actual, expected, name)

    def test_vendored_license_matches_third_party_record(self) -> None:
        vendored = (VENDOR / "LICENSE").read_bytes()
        recorded = (REPO_ROOT / "third_party" / "licenses" / "patent-disclosure-MIT.txt").read_bytes()
        self.assertEqual(vendored, recorded)

    def test_full_module_is_present(self) -> None:
        from oh_my_patent.render._vendor.patent_application import DEFAULT_STEMS

        self.assertIn("说明书", DEFAULT_STEMS)
        for name in ("md_to_docx.py", "math_to_omml.py", "latex_delimiters.py", "stdio_utf8.py"):
            self.assertTrue((VENDOR / name).is_file(), name)


class ProjectionTests(unittest.TestCase):
    def test_parts_and_section_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            parts = markdown_parts(_draft(Path(tmp), drawings=True), RenderOptions())
        self.assertEqual(
            set(parts), {"claims", "description", "abstract", "drawings"}
        )
        description = parts["description"]
        positions = [
            description.index(f"【{title}】")
            for title in ("技术领域", "背景技术", "发明内容", "附图说明", "具体实施方式")
        ]
        self.assertEqual(positions, sorted(positions))
        self.assertIn("## 一种缓冲队列处理方法", description)
        self.assertTrue(parts["claims"].startswith("# 权利要求书\n\n1. 一种缓冲队列"))
        self.assertIn("标记ABSTRACT_E", parts["abstract"])
        self.assertIn("![图1](", parts["drawings"])
        self.assertIn("图1", parts["drawings"])

    def test_split_projection_includes_abstract_drawing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            parts = markdown_parts(_draft(Path(tmp), drawings=True), RenderOptions(), split=True)
        self.assertIn("abstract_drawing", parts)
        self.assertIn("摘要附图", parts["abstract_drawing"])
        self.assertNotIn("![摘要附图]", parts["abstract"])


class RenderTests(unittest.TestCase):
    def test_boundary_layout_and_numbering(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "out.docx"
            render_docx_upstream(_draft(root), target, RenderOptions())

            document = Document(str(target))
            section = document.sections[0]
            self.assertAlmostEqual(section.page_width.mm, 210, delta=0.5)
            self.assertAlmostEqual(section.page_height.mm, 297, delta=0.5)
            self.assertAlmostEqual(section.top_margin.mm, 25, delta=0.5)
            self.assertAlmostEqual(section.left_margin.mm, 25, delta=0.5)
            self.assertAlmostEqual(section.right_margin.mm, 15, delta=0.5)
            self.assertAlmostEqual(section.bottom_margin.mm, 15, delta=0.5)

            heading1 = document.styles["Heading 1"]
            self.assertAlmostEqual(heading1.font.size.pt, TITLE_FONT_SIZE_PT)
            self.assertEqual(str(heading1.font.color.rgb), "000000")
            heading2 = document.styles["Heading 2"]
            self.assertAlmostEqual(heading2.font.size.pt, HEADING_NAME_FONT_SIZE_PT)
            self.assertEqual(str(heading2.font.color.rgb), "000000")

            normal = document.styles["Normal"]
            self.assertAlmostEqual(normal.font.size.pt, 12)
            rfonts = normal.element.rPr.rFonts
            self.assertEqual(rfonts.get(qn("w:eastAsia")), "宋体")
            self.assertEqual(rfonts.get(qn("w:hint")), "eastAsia")

            claim_paragraphs = [p for p in document.paragraphs if p.text.startswith("1. 一种")]
            self.assertEqual(len(claim_paragraphs), 1)
            claim = claim_paragraphs[0]
            self.assertEqual(claim.style.name, "Normal")
            self.assertIsNone(claim._element.find(".//" + qn("w:numPr")))

            body = next(p for p in document.paragraphs if "标记FIELD_A" in p.text)
            self.assertGreaterEqual(body.paragraph_format.first_line_indent.pt, 23.5)
            self.assertAlmostEqual(body.paragraph_format.line_spacing, 1.5)

    def test_markdown_table_keeps_data_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            draft = _draft(root)
            draft.embodiments += (
                "\n\n| 参数 | 定义 | 单位 |\n| --- | --- | --- |\n"
                "| $n$ | 样本数量PARAM_N | 个 |\n| $x_i$ | 第i次时延PARAM_X | ms |"
            )
            target = root / "table.docx"
            render_docx_upstream(draft, target, RenderOptions())
            document = Document(str(target))
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(len(document.tables[0].rows), 3)
            cells = [cell.text for row in document.tables[0].rows for cell in row.cells]
            self.assertIn("样本数量PARAM_N", cells)
            self.assertIn("第i次时延PARAM_X", cells)
            self.assertFalse(
                any(p.text.strip().startswith("|") for p in document.paragraphs),
                "表格数据行不应退化成带竖线的正文段落",
            )

    def test_math_becomes_native_objects(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "math.docx"
            render_docx_upstream(_draft(root, math=True), target, RenderOptions())
            document = Document(str(target))
            equations = document.element.body.findall(".//" + M_OMML)
            self.assertGreaterEqual(len(equations), 2)
            self.assertNotIn("$", "".join(p.text for p in document.paragraphs))

    def test_single_output_contains_all_parts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "all.docx"
            render_docx_upstream(_draft(root, drawings=True), target, RenderOptions())
            document = Document(str(target))
            titles = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
            self.assertEqual(titles, ["权利要求书", "说明书", "说明书摘要", "说明书附图"])

    def test_split_output_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            out = root / "split"
            written = render_docx_upstream(_draft(root, drawings=True), out, RenderOptions(), split=True)
            names = {path.name for path in written}
            self.assertEqual(
                names,
                {
                    "一种缓冲队列处理方法_权利要求书.docx",
                    "一种缓冲队列处理方法_说明书.docx",
                    "一种缓冲队列处理方法_说明书摘要.docx",
                    "一种缓冲队列处理方法_说明书附图.docx",
                    "一种缓冲队列处理方法_摘要附图.docx",
                },
            )
            for path in written:
                self.assertGreater(path.stat().st_size, 0)


class FailureGateTests(unittest.TestCase):
    def test_math_fallback_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            draft = _draft(root)
            draft.technical_field = "含无法转换的公式 $\\begin{cases}$。"
            with self.assertRaises(UpstreamRenderError) as ctx:
                render_docx_upstream(draft, root / "bad.docx", RenderOptions())
            self.assertIn("公式", str(ctx.exception))

    def test_bare_dollar_pair_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            draft = _draft(root)
            draft.technical_field = "行内写成 $$x +$$ 时不会被转换。"
            with self.assertRaises(UpstreamRenderError):
                render_docx_upstream(draft, root / "bad.docx", RenderOptions())

    def test_missing_image_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            draft = _draft(root)
            draft.drawings = [Drawing(number=1, caption="缺失图", path=str(root / "no-such.png"))]
            with self.assertRaises(UpstreamRenderError) as ctx:
                render_docx_upstream(draft, root / "bad.docx", RenderOptions())
            self.assertIn("图片", str(ctx.exception))

    def test_bare_paren_latex_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            draft = _draft(root)
            draft.technical_field = "均值 (M_{\\mathrm{total}}) 按普通括号书写。"
            with self.assertRaises(UpstreamRenderError) as ctx:
                render_docx_upstream(draft, root / "bad.docx", RenderOptions())
            self.assertIn("LaTeX", str(ctx.exception))

    def test_design_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            draft = PatentDraft(title="一种外观", patent_type=PatentType.DESIGN)
            with self.assertRaises(UpstreamRenderError) as ctx:
                render_docx_upstream(draft, root / "design.docx", RenderOptions())
            self.assertIn("外观设计", str(ctx.exception))


class CliTests(unittest.TestCase):
    DRAFT = """---
title: 一种测试方法
type: invention
---

# 权利要求书

1. 一种测试方法，其特征在于，包括接收记录、保存记录和输出记录。

# 说明书

本发明涉及数据处理领域，标记FIELD_A。

【技术领域】

本发明涉及数据处理领域，标记FIELD_A。

【背景技术】

输入记录按既定顺序读取，标记BACKGROUND_B。

【发明内容】

接收模块读取记录，缓冲模块保存记录。

【具体实施方式】

步骤S101，接收带序号的记录，标记END_D。

# 摘要

本合成示例仅用于测试，标记ABSTRACT_E。
"""

    def test_cli_build_with_upstream_renderer(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            draft_path = root / "draft.md"
            draft_path.write_text(self.DRAFT, encoding="utf-8")
            output = root / "out.docx"
            code = cli_main(
                ["build", str(draft_path), "-o", str(output), "--renderer", "upstream"]
            )
            self.assertEqual(code, 0)
            self.assertTrue(output.is_file())
            document = Document(str(output))
            titles = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
            self.assertEqual(titles, ["权利要求书", "说明书", "说明书摘要"])


if __name__ == "__main__":
    unittest.main()
