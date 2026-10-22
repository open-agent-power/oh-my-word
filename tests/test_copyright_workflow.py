"""Exercise the vendored stages, their local safety hooks and rendered manuals."""
import json
from pathlib import Path
import subprocess
import sys

from docx import Document
from PIL import Image
import pytest

from word_skills import copyright
from word_skills.common import read_json, write_json, sha256
from word_skills.copyright_package import prepare, apply_selection, package
from word_skills.copyright_workflow import VENDOR
from word_skills.copyright_manual import markdown_blocks


def stage(name, *args):
    return subprocess.run([sys.executable, "-m", "word_skills", "copyright", "workflow", name,
                           *map(str, args)], capture_output=True, text=True, encoding="utf-8")


def test_pinned_vendor_bytes():
    record = read_json(VENDOR / "UPSTREAM.json")
    assert record["version"] == "2.6"
    assert all(sha256(VENDOR / path) == checksum for path, checksum in record["files"].items())
    assert "Permission is hereby granted" in (VENDOR / "LICENSE").read_text()


def test_prepare_reuses_discovery_and_records_unreadable_design(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "logic.unknownlang").write_text("procedure run\n  return true\n")
    (source / "private.py").write_text("GITHUB_TOKEN='this_is_a_fixture_secret'\n")
    (source / ".env").write_text("TOKEN=fixture\n")
    (source / "README.md").write_text("# Example\nGITHUB_TOKEN='this_is_a_fixture_secret'\n")
    (source / "设计.wps").write_bytes(b"legacy document fixture")
    case = tmp_path / "case"
    prepare(source, case, name="测试软件", version="V1.0")
    manifest = read_json(case / "source-manifest.json")
    assert [item["path"] for item in manifest["files"]] == ["logic.unknownlang"]
    candidates = read_json(case / "workflow/草稿/代码文件选择.json")
    assert [item["path"] for item in candidates["files"]] == ["logic.unknownlang"]
    assert not any(item["selected"] for item in candidates["files"])
    analysis = read_json(case / "workflow/analysis/project.json")
    assert "this_is_a_fixture_secret" not in json.dumps(analysis)
    documents = read_json(case / "workflow/草稿/业务理解证据.json")["documents"]
    legacy = next(item for item in documents if item["path"] == "设计.wps")
    assert not legacy["text_extracted"] and "WPS" in legacy["extraction_note"]


def test_changed_business_cannot_reuse_confirmation(tmp_path):
    work = tmp_path / "workflow"
    write_json(work / "草稿/业务理解.json", {"confirmation_required": True, "purpose": "fixture"})
    result = stage("confirm", "--workdir", work, "--stage", "business", "--note", "fixture owner confirmation")
    assert result.returncode == 0, result.stderr
    context = read_json(work / "草稿/业务理解.json")
    context["purpose"] = "changed after confirmation"
    write_json(work / "草稿/业务理解.json", context)
    write_json(work / "analysis.json", {})
    result = stage("manual", "--analysis", work / "analysis.json", "--software-name", "测试软件",
                   "--business-context", work / "草稿/业务理解.json", "--out-dir", work / "manual")
    assert result.returncode != 0 and "stale" in result.stderr
    assert not (work / "manual/操作手册.md").exists()


def test_fields_do_not_turn_test_machine_into_development_facts(tmp_path):
    write_json(tmp_path / "analysis.json", {"project_root": str(tmp_path), "source": {"total_line_count": 2}})
    write_json(tmp_path / "business.json", {"confirmation_required": True, "user_confirmed": False,
               "application_purpose": "fixture purpose", "main_functions": "功能说明" * 130,
               "technical_characteristics": "真实项目字段"})
    write_json(tmp_path / "code.json", {"submitted_page_count": 1, "mode": "all_under_60_pages"})
    result = stage("fields", "--preview", "--analysis", tmp_path / "analysis.json",
                   "--code-manifest", tmp_path / "code.json", "--software-name", "测试软件",
                   "--business-context", tmp_path / "business.json", "--out-dir", tmp_path / "draft")
    assert result.returncode == 0, result.stderr
    text = (tmp_path / "draft/申请表信息.md").read_text(encoding="utf-8")
    assert "➤软件开发环境 / 开发工具：待用户确认" in text
    assert "➤开发方式：待用户确认" in text
    assert "➤运行的硬件环境：待用户确认" in text
    assert not read_json(tmp_path / "business.json")["user_confirmed"]


@pytest.mark.parametrize("authored,answer,expected", [
    ("按分隔符截断输入，返回处理后的字符串。", None, "按分隔符截断输入，返回处理后的字符串。"),
    ("", None, "待用户确认"),
    ("业务模型中的原始说明", "仅执行命令并打印结果。", "仅执行命令并打印结果。"),
    ("事实说明" * 400, None, "事实说明" * 400),
])
def test_fields_never_pad_or_replace_authored_functionality(tmp_path, authored, answer, expected):
    write_json(tmp_path / "analysis.json", {"project_root": str(tmp_path), "source": {"total_line_count": 2}})
    write_json(tmp_path / "business.json", {"confirmation_required": True, "user_confirmed": False,
               "main_functions": authored, "application_purpose": "文本处理"})
    write_json(tmp_path / "code.json", {"submitted_page_count": 1, "mode": "all_under_60_pages"})
    write_json(tmp_path / "answers.json", {"软件的主要功能": answer} if answer else {})
    result = stage("fields", "--preview", "--analysis", tmp_path / "analysis.json",
                   "--code-manifest", tmp_path / "code.json", "--software-name", "测试软件",
                   "--business-context", tmp_path / "business.json", "--answers", tmp_path / "answers.json",
                   "--out-dir", tmp_path / "draft")
    assert result.returncode == 0, result.stderr
    text = (tmp_path / "draft/申请表信息.md").read_text(encoding="utf-8")
    row = next(line for line in text.splitlines() if line.startswith("➤软件的主要功能："))
    assert row == "➤软件的主要功能：" + expected
    assert "系统提供数据持久化存储" not in text
    assert "固定上游兼容校验" in text
    assert read_json(tmp_path / "business.json")["user_confirmed"] is False


@pytest.fixture
def manual_case(tmp_path):
    root, case = tmp_path / "source", tmp_path / "case"
    root.mkdir()
    (root / "main.py").write_text("print('hello')\n")
    prepare(root, case, name="测试软件", version="V1.0")
    selection = read_json(case / "selection.json")
    selection["reviewed_by"] = "fixture reviewer"
    selection["files"][0].update(selected=True, reason="real fixture entry")
    write_json(case / "selection.json", selection)
    apply_selection(case / "copyright.json")
    Image.new("RGB", (100, 40), "white").save(case / "result.png")
    (case / "manual.md").write_text(
        "# 测试软件 V1.0\n\n## 查询结果\n\n| 项目 | 结果 |\n| --- | --- |\n| 查询 | hello |\n\n"
        "```text\noks recall hello\nhello\n```\n\n![查询输出](result.png)\n", encoding="utf-8")
    value = read_json(case / "copyright.json")
    value["application"]["category"] = "命令行软件"
    value["manual"]["verified_against"] = "fixture source and output"
    value["screenshots"] = [{"path": "result.png", "feature": "查询输出"}]
    write_json(case / "copyright.json", value)
    return case


def test_manual_tables_commands_and_images_reach_word(manual_case, tmp_path):
    result = package(manual_case / "copyright.json", tmp_path / "delivery", draft=True, no_preview=True)
    assert result["state"] == "draft"
    folder = tmp_path / "delivery/manual"
    doc = Document(folder / "output.docx")
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 1
    assert any("oks recall hello\nhello" == p.text for p in doc.paragraphs)
    draft = read_json(folder / "audit/draft.json")
    image = next(b for s in draft["sections"] for b in s["blocks"] if b["type"] == "image")
    assert sha256(folder / "audit" / image["path"]) == image["sha256"]
    (folder / "audit" / image["path"]).write_bytes(b"changed image fixture")
    from word_skills.evidence import validate
    assert any(item["code"] == "IMAGE_SOURCE" for item in validate(draft, folder / "audit")["issues"])


def test_undeclared_manual_image_is_rejected(manual_case, tmp_path):
    (manual_case / "manual.md").write_text("测试软件 V1.0\n\n![非声明图片](other.png)", encoding="utf-8")
    with pytest.raises(ValueError, match="declared"):
        copyright.build(manual_case / "copyright.json", tmp_path / "bad", purpose="preview", document="manual", no_preview=True)


def test_image_inside_table_is_explicitly_rejected(manual_case, tmp_path):
    (manual_case / "manual.md").write_text(
        "测试软件 V1.0\n\n| 输出 | 说明 |\n| --- | --- |\n| ![结果](result.png) | 查询输出 |\n", encoding="utf-8")
    with pytest.raises(ValueError, match="standalone"):
        copyright.build(manual_case / "copyright.json", tmp_path / "bad-table", purpose="preview", document="manual", no_preview=True)


def test_manual_lists_and_nested_commands_reach_word(manual_case, tmp_path):
    (manual_case / "manual.md").write_text(
        "# 测试软件 V1.0\n\n3. 启动服务\n   - 检查输出\n\n"
        "     ```sh\n     oks recall hello\n     ```\n\n4. 打开控制台\n",
        encoding="utf-8")
    out = tmp_path / "numbered-manual"
    result = copyright.build(manual_case / "copyright.json", out,
                             purpose="preview", document="manual", no_preview=True)
    assert result["state"] == "draft"
    texts = [p.text for p in Document(out / "output.docx").paragraphs]
    assert "3. 启动服务" in texts and "4. 打开控制台" in texts
    assert "  - 检查输出" in texts and "oks recall hello" in texts


def test_list_markers_do_not_invent_numbers_or_modify_commands(tmp_path):
    blocks = markdown_blocks("1. First\n1. Second\n\n   ```sh\n   run --port 8080\n   ```\n", [], tmp_path, tmp_path)
    assert [b["text"] for b in blocks] == ["1. First", "1. Second", "run --port 8080"]


@pytest.mark.parametrize("text,expected", [
    ("1. - Nested\n", ["1.", "  - Nested"]),
    ("3) Third\n4) Fourth\n", ["3) Third", "4) Fourth"]),
    ("> 3. Third\n> 4. Fourth\n", ["3. Third", "4. Fourth"]),
])
def test_list_marker_variants(tmp_path, text, expected):
    assert [b["text"] for b in markdown_blocks(text, [], tmp_path, tmp_path)] == expected


@pytest.mark.parametrize("html", [
    "<table><tr><td>Critical requirement</td></tr></table>",
    "Press <kbd>Enter</kbd> to continue.",
    "| Action | Result |\n| --- | --- |\n| Run | <b>Required</b> |",
])
def test_unsupported_manual_html_fails_before_word_output(manual_case, tmp_path, html):
    (manual_case / "manual.md").write_text("# 测试软件 V1.0\n\n" + html + "\n", encoding="utf-8")
    out = tmp_path / "html-manual"
    with pytest.raises(ValueError, match="Raw HTML.*line.*Markdown"):
        copyright.build(manual_case / "copyright.json", out,
                        purpose="preview", document="manual", no_preview=True)
    assert not (out / "output.docx").exists()


def test_workflow_selection_uses_same_files_and_rejects_changed_source(manual_case):
    path = manual_case / "copyright.json"
    original = read_json(path)["code_material"]["files"][0]
    selection_path = manual_case / "workflow-selection.json"
    write_json(selection_path, {"files": [{"path": original["path"], "selected": True, "model_reason": "actual fixture entry"}]})
    apply_selection(path, workflow_selection=selection_path, reviewed_by="fixture reviewer")
    assert read_json(path)["code_material"]["files"][0]["sha256"] == original["sha256"]
    root = Path(read_json(manual_case / "source-manifest.json")["source_root"])
    (root / original["path"]).write_text("changed\n")
    with pytest.raises(ValueError, match="changed"):
        apply_selection(path, workflow_selection=selection_path, reviewed_by="fixture reviewer")


def test_unknown_source_extension_survives_full_build(tmp_path):
    source = tmp_path / "src"
    source.mkdir()
    (source / "main.unknownlang").write_text("procedure start\n  return true\n")
    case = tmp_path / "case"
    prepare(source, case, name="测试软件", version="V1.0")
    selection = read_json(case / "selection.json")
    selection["reviewed_by"] = "fixture reviewer"
    selection["files"][0].update(selected=True, reason="actual fixture entry")
    write_json(case / "selection.json", selection)
    apply_selection(case / "copyright.json")
    value = read_json(case / "copyright.json")
    value["application"]["category"] = "命令行软件"
    value["manual"]["verified_against"] = "fixture source"
    (case / "manual.md").write_text("测试软件 V1.0\n启动后返回true。", encoding="utf-8")
    write_json(case / "copyright.json", value)
    result = package(case / "copyright.json", tmp_path / "delivery", draft=True, no_preview=True)
    assert result["state"] == "draft"
    doc = Document(tmp_path / "delivery/source/output.docx")
    assert any(p.text == "  return true" for p in doc.paragraphs)
