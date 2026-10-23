from docx import Document
from docx.oxml.ns import qn

from word_skills.render import append_content, default_document, render


def _draft(blocks):
    return {
        "title": "版式测试",
        "sections": [{"title": "正文", "blocks": blocks}],
    }


def test_table_rows_are_marked_unsplittable_and_header_repeats():
    doc = default_document()
    append_content(doc, _draft([
        {"type": "table", "headers": ["参数", "说明"],
         "rows": [["采样率", "每秒采集一次"]]},
    ]))

    table = doc.tables[0]
    assert table.cell(1, 0).text == "采样率"
    assert table.cell(1, 1).text == "每秒采集一次"
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in table.rows)
    # These OOXML properties request the behavior; only a rendered PDF can
    # establish how a particular Word/LibreOffice version paginates a long row.


def test_table_has_no_trailing_spacer_paragraph():
    doc = default_document()
    append_content(doc, _draft([
        {"type": "table", "headers": ["列"], "rows": [["值"]]},
    ]))

    body_children = list(doc.element.body)
    content_children = [node for node in body_children if node.tag != qn("w:sectPr")]
    assert content_children[-1].tag == qn("w:tbl")


def test_plain_paragraph_and_code_blocks_keep_backticks_verbatim():
    plain = "普通文本中的 `token` 保持原样"
    code = "print(`source backticks`)"
    doc = default_document()
    append_content(doc, _draft([
        {"type": "paragraph", "text": plain},
        {"type": "code", "text": code},
        {"type": "table", "headers": ["示例"], "rows": [["`table token`"]]},
    ]))

    assert plain in [p.text for p in doc.paragraphs]
    assert code in [p.text for p in doc.paragraphs]
    assert doc.tables[0].cell(1, 0).text == "`table token`"


def test_markdown_inline_code_is_opt_in_and_keeps_visible_text():
    doc = default_document()
    append_content(doc, _draft([
        {"type": "paragraph", "text": "运行 `python app.py` 即可。",
         "inline_format": "markdown"},
    ]))

    paragraph = doc.paragraphs[-1]
    assert paragraph.text == "运行 python app.py 即可。"
    code_runs = [run for run in paragraph.runs if run.font.name == "Consolas"]
    assert [run.text for run in code_runs] == ["python app.py"]


def test_default_text_and_linked_heading_colors_are_fixed_black(tmp_path):
    from zipfile import ZipFile
    from lxml import etree

    draft = {"schema_version": 1, "mode": "copyright-manual", "title": "操作说明",
             "software_name": "示例", "software_version": "0.3.0",
             "sections": [{"id": "s", "title": "使用方法", "blocks": [
                 {"id": "p", "type": "paragraph", "text": "运行命令", "intent": "analysis", "claim_ids": []},
                 {"id": "c", "type": "code", "text": "print('hello')", "intent": "analysis", "claim_ids": []}]}]}
    target = tmp_path / "default.docx"
    render(draft, target)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(target) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
        for sid in ("Normal", "Title", "TitleChar", "Heading1", "Heading1Char", "Heading2Char", "DefaultParagraphFont"):
            colors = styles.xpath(f'//w:style[@w:styleId="{sid}"]/w:rPr/w:color', namespaces=ns)
            assert len(colors) == 1 and colors[0].attrib == {qn("w:val"): "000000"}
        assert styles.xpath('//w:docDefaults/w:rPrDefault/w:rPr/w:color/@w:val', namespaces=ns) == ["000000"]
        for name in archive.namelist():
            if name == "word/document.xml" or name.startswith(("word/header", "word/footer")) and name.endswith(".xml"):
                story = etree.fromstring(archive.read(name))
                for run in story.xpath('//w:r', namespaces=ns):
                    colors = run.xpath('./w:rPr/w:color', namespaces=ns)
                    assert len(colors) == 1 and colors[0].attrib == {qn("w:val"): "000000"}
    assert "print('hello')" in [p.text for p in Document(target).paragraphs]


def test_customer_template_font_color_is_preserved(tmp_path):
    from docx.shared import RGBColor

    template = Document()
    template.styles["Title"].font.color.rgb = RGBColor.from_string("AA1122")
    reference = tmp_path / "reference.docx"
    template.save(reference)
    target = tmp_path / "output.docx"
    draft = {"schema_version": 1, "mode": "copyright-manual", "title": "标题",
             "sections": [{"id": "s", "title": "正文", "blocks": [
                 {"id": "p", "type": "paragraph", "text": "内容", "intent": "analysis", "claim_ids": []}]}]}
    render(draft, target, template=reference)
    assert Document(target).styles["Title"].font.color.rgb == RGBColor.from_string("AA1122")
