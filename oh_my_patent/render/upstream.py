"""专利 DOCX 制作的上游复用适配层。

固定上游：``handsomestWei/patent-disclosure-skill`` 提交
``6935a14d342ed9c7a07440c999d1b7bf1852a701``（MIT），原件逐字节保存在
``_vendor/patent_application/``，哈希见同目录 ``UPSTREAM.json``。

本模块只做三件事，不重写上游的 Markdown 解析、表格或公式转换：

1. 把 :class:`~oh_my_patent.schema.PatentDraft` 投影成上游认识的
   「权利要求书 / 说明书 / 说明书摘要 / 说明书附图」Markdown；
2. 调用上游 ``convert_md_to_docx`` 生成文档；
3. 在边界套用本项目的版面配置（A4、黑色标题与正文、宋体、默认 1.25 倍行距
   ——实测行间空白 2.96mm，落在规定 2.5~3.5mm 内、2 字符首行缩进、页边距），
   并把上游「缺图/坏公式仍返回成功」的行为拦成明确失败。

合法输入仍以现有稿件模型、外层审阅与交接门禁为准；这里的失败只表示
「本次 DOCX 制作不可交付」，不替代 ``lint``、``final`` 或人工复核。
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .. import oxml
from ..schema import PatentDraft, PatentType
from ..spec import (
    HEADING_NAME_FONT_SIZE_PT,
    PART_TITLES,
    SECTION_FONT_SIZE_PT,
    SECTION_TITLE_PATTERN,
    TITLE_FONT_SIZE_PT,
    RenderOptions,
)
from ._vendor.patent_application import (
    convert_md_to_docx,
    find_bare_paren_latex,
    format_hits_report,
    get_math_stats,
)
from .docx import DocxRenderer, _safe_filename

__all__ = [
    "UpstreamRenderError",
    "build_upstream_documents",
    "markdown_parts",
    "render_docx_upstream",
    "upstream_identity",
]

_VENDOR_DIR = Path(__file__).resolve().parent / "_vendor" / "patent_application"
_MANIFEST = _VENDOR_DIR / "UPSTREAM.json"
_EXPECTED_COMMIT = "6935a14d342ed9c7a07440c999d1b7bf1852a701"

_W_DRAWING = qn("w:drawing")
_M_OMML = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"

#: 上游把解析失败的图片写成这些正文标记，而不是抛错。
_MISSING_MARKERS = ("[图片缺失:", "[图片无法嵌入:", "[行内公式图缺失:")

#: 公式分隔符。与上游 ``_maybe_render_math_md`` 的探测口径保持同一范围：
#: 行内 ``$...$``、块级 ``$$...$$`` 与 ``\[...\]``。
_MATH_TOKEN_RE = re.compile(r"\$\$?|\\\[")

#: 说明书附图下的居中图号行，如 ``图1``。
_CAPTION_RE = re.compile(r"图\d+")

#: 各部分渲染顺序；与 :data:`oh_my_patent.spec.PART_ORDER` 一致并允许摘要附图。
_PART_SEQUENCE = ("claims", "description", "abstract", "drawings", "abstract_drawing")


class UpstreamRenderError(ValueError):
    """上游制作在当前输入下不可交付；调用方不得把它当成生成成功。"""


# --------------------------------------------------------------------------
# 固定身份
# --------------------------------------------------------------------------

def upstream_identity(*, verify: bool = True) -> dict:
    """读取固定原件清单；``verify`` 时逐个校验 SHA256。

    运行时校验是「实际执行的是固定上游代码」的证据：哈希不符即拒绝制作，
    避免有人就地打补丁后仍声称复用原版。
    """
    manifest = json.loads(_MANIFEST.read_text(encoding="utf-8"))
    if manifest.get("commit") != _EXPECTED_COMMIT:
        raise UpstreamRenderError(
            f"上游清单提交 {manifest.get('commit')} 与本适配层固定的 {_EXPECTED_COMMIT} 不一致"
        )
    if not verify:
        return manifest
    mismatched: list[str] = []
    for name, expected in manifest["files"].items():
        path = _VENDOR_DIR / name
        if not path.is_file():
            mismatched.append(f"{name}: 缺失")
            continue
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != expected:
            mismatched.append(f"{name}: {actual} != {expected}")
    if mismatched:
        raise UpstreamRenderError(
            "上游原件 SHA256 与 UPSTREAM.json 不符，拒绝制作：" + "; ".join(mismatched)
        )
    return manifest


# --------------------------------------------------------------------------
# 稿件 → 上游 Markdown
# --------------------------------------------------------------------------

def markdown_parts(
    draft: PatentDraft,
    options: RenderOptions | None = None,
    *,
    split: bool = False,
) -> dict[str, str]:
    """把稿件投影成上游认识的 Markdown 文书（键为部分名）。"""
    if draft.patent_type is PatentType.DESIGN:
        raise UpstreamRenderError(
            "外观设计暂未接入上游四文书转换器；请使用本地渲染器（默认 --renderer local）。"
        )
    options = options or RenderOptions()
    renderer = DocxRenderer(draft, options)

    parts: dict[str, str] = {}

    claims = [claim.render_text().strip() for claim in draft.claims if claim.text.strip()]
    if claims:
        lines = [f"# {PART_TITLES['claims']}"] if options.include_part_titles else []
        parts["claims"] = _join(lines + claims)

    description = _description_markdown(draft, renderer, options)
    if description:
        parts["description"] = description

    abstract = _abstract_markdown(draft, options, split=split)
    if abstract:
        parts["abstract"] = abstract

    drawings = _drawings_markdown(draft, options)
    if drawings:
        parts["drawings"] = drawings

    figure = draft.abstract_figure
    if split and figure is not None and figure.path:
        lines = [f"# {PART_TITLES['abstract_drawing']}"] if options.include_part_titles else []
        parts["abstract_drawing"] = _join(lines + [f"![摘要附图]({_image_ref(figure.path)})"])

    return parts


def _description_markdown(
    draft: PatentDraft, renderer: DocxRenderer, options: RenderOptions
) -> str:
    """说明书：发明名称 + 法定章节。章节组装复用本地渲染层的同一实现。

    直接使用 ``DocxRenderer._description_blocks`` 是有意的：正文如何切段、
    「发明内容」如何按截面样式展开，只应有一处定义，避免两套后端各自漂移。
    """
    blocks = renderer._description_blocks()
    if not blocks:
        return ""

    lines: list[str] = []
    if options.include_part_titles:
        lines.append(f"# {PART_TITLES['description']}")
    if draft.title.strip():
        lines.append(f"## {draft.title.strip()}")
    for block in blocks:
        if block.level == 1:
            lines.append(f"### {SECTION_TITLE_PATTERN.format(title=block.title)}")
        else:
            lines.append(f"#### {block.title}")
        lines.extend(block.paragraphs)
    return _join(lines)


def _abstract_markdown(
    draft: PatentDraft, options: RenderOptions, *, split: bool
) -> str:
    paragraphs = DocxRenderer._paragraphs_of(draft.abstract)
    figure = draft.abstract_figure
    content = list(paragraphs)
    if not split and figure is not None:
        if not figure.path:
            raise UpstreamRenderError("摘要附图已登记但缺少图片文件，拒绝生成不完整文书。")
        content.append(f"![摘要附图]({_image_ref(figure.path)})")
    if not content:
        return ""
    lines = [f"# {PART_TITLES['abstract']}"] if options.include_part_titles else []
    return _join(lines + content)


def _drawings_markdown(draft: PatentDraft, options: RenderOptions) -> str:
    drawings = sorted(draft.drawings, key=lambda item: item.number)
    if not drawings:
        return ""
    lines = [f"# {PART_TITLES['drawings']}"] if options.include_part_titles else []
    for drawing in drawings:
        if not drawing.path:
            raise UpstreamRenderError(
                f"图{drawing.number}已登记但缺少图片文件，拒绝生成不完整文书。"
            )
        lines.append(f"![图{drawing.number}]({_image_ref(drawing.path)})")
        lines.append(f"图{drawing.number}")
    return _join(lines)


def _join(lines: list[str]) -> str:
    """拼接上游 Markdown，保持表格行紧凑，其余段落之间留空行。

    上游按「行」解析表格，表格行之间一旦插入空行，后续行就会退化成正文
    （实测：表头进了 Word 表格，数据行变成带 ``|`` 的普通段落）。因此连续
    的 ``|…|`` 行必须用单换行连接。
    """
    chunks: list[str] = []
    table: list[str] = []
    for line in lines:
        if line.strip().startswith("|"):
            table.append(line.strip())
            continue
        if table:
            chunks.append("\n".join(table))
            table = []
        if line.strip():
            chunks.append(line.strip())
    if table:
        chunks.append("\n".join(table))
    return "\n\n".join(chunk for chunk in chunks if chunk.strip())


def _image_ref(path: str) -> str:
    return str(Path(path).resolve()).replace("\\", "/")


# --------------------------------------------------------------------------
# 制作
# --------------------------------------------------------------------------

def build_upstream_documents(
    draft: PatentDraft,
    options: RenderOptions | None = None,
    *,
    split: bool = False,
) -> list[tuple[str, Document]]:
    """生成各部分文档；任何「上游成功但其实不可交付」的情况都抛错。"""
    options = options or RenderOptions()
    options.validate()
    upstream_identity(verify=True)
    parts = markdown_parts(draft, options, split=split)
    if not parts:
        raise UpstreamRenderError("稿件没有可制作的内容。")
    return [
        (key, _convert_part(key, parts[key], options))
        for key in _PART_SEQUENCE
        if key in parts
    ]


def _convert_part(key: str, markdown: str, options: RenderOptions) -> Document:
    title = PART_TITLES.get(key, key)

    hits = find_bare_paren_latex(markdown)
    if hits:
        raise UpstreamRenderError(
            f"{title}含普通括号 LaTeX 写法（{len(hits)} 处），Word 会当纯文本，拒绝转换。"
            f"\n{format_hits_report(hits)}"
        )

    has_math = bool(_MATH_TOKEN_RE.search(markdown))
    expected_images = sum(
        1 for line in markdown.splitlines() if line.lstrip().startswith("![")
    )
    try:
        # 上游把 base_dir 当作「能否嵌入图片」的总开关（为 None 时不嵌图且不报错），
        # 因此必须传入一个真实目录；图片引用本身已绝对化，不依赖相对解析。
        document = convert_md_to_docx(markdown, Path.cwd(), prefer_omml=True)
    except Exception as exc:  # 上游异常统一转成本层可识别的失败
        raise UpstreamRenderError(f"{title}上游转换失败：{exc}") from exc

    stats = get_math_stats()
    if stats.text:
        raise UpstreamRenderError(
            f"{title}有 {stats.text} 个公式回退为 LaTeX 原文，拒绝当作制作成功："
            + "；".join(stats.text_latex[:3])
        )
    if has_math and stats.omml + stats.png == 0:
        raise UpstreamRenderError(
            f"{title}含公式但没有任何公式被转换为对象，请检查 LaTeX 与依赖（latex2mathml）。"
        )

    text = _document_text(document)
    for marker in _MISSING_MARKERS:
        if marker in text:
            raise UpstreamRenderError(
                f"{title}存在未嵌入的图片/公式图（{marker}），拒绝生成不完整文书。"
            )

    embedded = len(document.element.body.findall(".//" + _W_DRAWING))
    if embedded < expected_images:
        raise UpstreamRenderError(
            f"{title}应嵌入 {expected_images} 张图片，实际 {embedded} 张，拒绝生成不完整文书。"
        )

    _apply_boundary_layout(document, options)
    return document


def render_docx_upstream(
    draft: PatentDraft,
    output: str | Path,
    options: RenderOptions | None = None,
    split: bool = False,
) -> list[Path]:
    """用上游转换器制作 .docx；``split`` 语义与本地渲染器一致。"""
    options = options or RenderOptions()
    documents = build_upstream_documents(draft, options, split=split)
    output = Path(output)

    if split:
        output.mkdir(parents=True, exist_ok=True)
        stem = _safe_filename(draft.title)
        written: list[Path] = []
        for key, document in documents:
            target = output / f"{stem}_{PART_TITLES.get(key, key)}.docx"
            document.save(str(target))
            written.append(target)
        return written

    output.parent.mkdir(parents=True, exist_ok=True)
    if len(documents) == 1:
        documents[0][1].save(str(output))
        return [output]
    _compose(documents, options, output)
    return [output]


def _compose(
    documents: list[tuple[str, Document]], options: RenderOptions, output: Path
) -> None:
    """把多份文书合成一个审阅文件；分页与本地渲染器保持一致。"""
    try:
        from docxcompose.composer import Composer
    except ImportError as exc:  # pragma: no cover - 依赖缺失时才触发
        raise UpstreamRenderError(
            "单文件输出需要 docxcompose；它已列为基础依赖，请重新安装后再试。"
        ) from exc

    composer = Composer(documents[0][1])
    for _, document in documents[1:]:
        if options.page_break_between_parts and document.paragraphs:
            oxml.set_page_break_before(document.paragraphs[0], True)
        composer.append(document)
    composer.save(str(output))


# --------------------------------------------------------------------------
# 边界版式
# --------------------------------------------------------------------------

_HEADING_LAYOUT: dict[str, tuple[float, bool, int]] = {
    "Heading 1": (TITLE_FONT_SIZE_PT, True, WD_ALIGN_PARAGRAPH.CENTER),
    "Heading 2": (HEADING_NAME_FONT_SIZE_PT, True, WD_ALIGN_PARAGRAPH.CENTER),
    "Heading 3": (SECTION_FONT_SIZE_PT, False, WD_ALIGN_PARAGRAPH.LEFT),
    "Heading 4": (SECTION_FONT_SIZE_PT, True, WD_ALIGN_PARAGRAPH.LEFT),
}

_LIST_STYLES = frozenset({"List Number", "List Bullet"})


def _restore_literal_numbering(document: Document) -> None:
    """把上游的自动编号列表还原成纯文本编号。

    上游把以 ``1. …`` / ``- …`` 开头的行解析成 Word 列表（``numPr``），本项目
    的既有排版明确不用自动编号：WPS 与 Office 对 ``numbering.xml`` 的实现不同，
    权利要求编号会在 WPS 里错位或显示成圆点。本地渲染器一直写纯文本编号，
    这里在边界还原成同样的文本，编号按各自列表实例从 1 重新计数。
    """
    counters: dict[str, int] = {}
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        p_pr = paragraph._element.pPr
        num_pr = p_pr.find(qn("w:numPr")) if p_pr is not None else None
        if num_pr is None and style_name not in _LIST_STYLES:
            continue

        num_id = "0"
        if num_pr is not None:
            num_id_element = num_pr.find(qn("w:numId"))
            if num_id_element is not None:
                num_id = num_id_element.get(qn("w:val")) or "0"
            p_pr.remove(num_pr)

        ordered = style_name == "List Number"
        paragraph.style = document.styles["Normal"]
        if ordered:
            counters[num_id] = counters.get(num_id, 0) + 1
            marker = f"{counters[num_id]}. "
        else:
            marker = "- "
        _prepend_run(paragraph, marker)


def _prepend_run(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    element = run._element
    paragraph._element.remove(element)
    first = next(
        (child for child in paragraph._element if child.tag != qn("w:pPr")), None
    )
    if first is None:
        paragraph._element.append(element)
    else:
        first.addprevious(element)


def _restore_table_pipes(document: Document) -> None:
    """还原表格单元格里的转义竖线。

    上游 ``_split_table_cells`` 会把 ``a \\| b`` 当作同一单元格内容，但把
    反斜杠原样留在文字里；Word 里应显示为 ``a | b``。这里只在表格单元格内
    还原，不影响正文的其它反斜杠。
    """
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if "\\|" in run.text:
                            run.text = run.text.replace("\\|", "|")


def _apply_boundary_layout(document: Document, options: RenderOptions) -> None:
    """把上游用 Word 默认模板生成的文档改成本项目要求的版面。

    上游没有版式配置入口（纸张、标题配色、字号、行距、缩进都取默认模板）。
    在边界统一收口，比给上游源码打补丁更容易核对：原件哈希保持可验证，
    差异只在下面这一处、且有测试覆盖。
    """
    oxml.configure_document(document, options)
    _restore_literal_numbering(document)
    _restore_table_pipes(document)
    western, east_asian = options.resolved_fonts()

    for name, (size, bold, alignment) in _HEADING_LAYOUT.items():
        _configure_heading(
            document, name, size, bold, alignment, western, east_asian, options.line_spacing
        )

    for paragraph in document.paragraphs:
        _format_paragraph(paragraph, options, western, east_asian)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    oxml.set_paragraph_format(
                        paragraph,
                        line_spacing=options.line_spacing,
                        first_line_indent_chars=0,
                        space_before_pt=0,
                        space_after_pt=0,
                    )
                    _set_runs(paragraph, options.body_size, western, east_asian)

    if options.page_numbers:
        oxml.add_page_number_footer(
            document.sections[0], western, east_asian, options.body_size
        )


def _configure_heading(
    document: Document,
    name: str,
    size_pt: float,
    bold: bool,
    alignment: int,
    western: str,
    east_asian: str,
    line_spacing: float,
) -> None:
    style = document.styles[name]
    style.font.name = western
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asian)
    rfonts.set(qn("w:cs"), western)
    rfonts.set(qn("w:hint"), "eastAsia")

    fmt = style.paragraph_format
    fmt.alignment = alignment
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = line_spacing


def _format_paragraph(
    paragraph, options: RenderOptions, western: str, east_asian: str
) -> None:
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    text = paragraph.text.strip()

    if style_name.startswith("Heading"):
        size, bold, alignment = _HEADING_LAYOUT.get(style_name, (SECTION_FONT_SIZE_PT, False, WD_ALIGN_PARAGRAPH.LEFT))
        paragraph.alignment = alignment
        oxml.set_paragraph_format(
            paragraph,
            line_spacing=options.line_spacing,
            first_line_indent_chars=0,
            space_before_pt=0,
            space_after_pt=0,
        )
        _set_runs(paragraph, size, western, east_asian)
        return

    if text:
        if _CAPTION_RE.fullmatch(text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            oxml.set_paragraph_format(
                paragraph,
                line_spacing=options.line_spacing,
                first_line_indent_chars=0,
                space_before_pt=0,
                space_after_pt=0,
            )
        else:
            has_math = bool(paragraph._element.findall(".//" + _M_OMML))
            oxml.set_paragraph_format(
                paragraph,
                line_spacing=options.line_spacing,
                first_line_indent_chars=options.first_line_indent,
                space_before_pt=0,
                space_after_pt=0,
                # 含公式的段落关闭中西文自动间距，否则 Word 会拆散 λ1、J(θ)。
                cjk_latin_spacing=not has_math,
            )
        _set_runs(paragraph, options.body_size, western, east_asian)
        return

    if paragraph._element.findall(".//" + _W_DRAWING):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        oxml.set_paragraph_format(
            paragraph,
            line_spacing=1.0,
            first_line_indent_chars=0,
            space_before_pt=0,
            space_after_pt=6,
        )


def _set_runs(paragraph, size_pt: float, western: str, east_asian: str) -> None:
    for run in paragraph.runs:
        if not run.text:
            continue
        oxml.set_run_font(run, western, east_asian, size_pt=size_pt)


def _document_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)
